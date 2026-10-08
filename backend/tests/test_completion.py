"""Regression checks use explicit test observations; none enter production data."""
import os
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["FREE_TIER_MODE"] = "1"
import unittest
from datetime import date, timedelta, datetime
from types import SimpleNamespace
from unittest.mock import patch
from io import BytesIO
import pandas as pd
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from openpyxl import load_workbook
from app.models import Base, Company, OHLCV, RSHistory
from app.api import market
from app.services import rs_universe_backfill as backfill
from app.services.rs_history import persist_history, decode_points
from app.services.providers.yahoo_provider import YahooProvider


def observations(symbol="AAPL", years=2):
    start = date.today() - timedelta(days=366 * years + 120)
    rows = []
    for i in range((date.today() - start).days + 1):
        day = start + timedelta(days=i)
        if day.weekday() < 5:
            close = 100 + i / 100 + (i % 19) / 10
            rows.append({"date": day, "open": close - .3, "high": close + 1, "low": close - 1, "close": close, "volume": 10000})
    return rows


class CompletionTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine)
        self.db = self.Session()
        self.db.add_all([Company(symbol="AAPL", exchange="US", name="Apple Inc.", is_active=1), Company(symbol="RELIANCE", exchange="NSE", name="Reliance Industries", isin="INE002A01018", is_active=1)])
        self.db.commit()
        market._LIVE_ROW_CACHE.clear()
        backfill._symbol_retry_after.clear()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def test_legacy_history_is_migrated_before_compaction(self):
        rows = observations()
        self.db.add_all([OHLCV(symbol="AAPL", exchange="US", **r) for r in rows])
        self.db.commit()
        selected = backfill._candidate_companies(self.db, "US", 5)
        self.assertIn("AAPL", [c.symbol for c in selected])
        persist_history(self.db, "AAPL", "US", rows, "yahoo")
        self.assertNotIn("AAPL", [c.symbol for c in backfill._candidate_companies(self.db, "US", 5)])

    def test_compact_history_has_bounded_storage(self):
        persist_history(self.db, "AAPL", "US", observations(years=20), "yahoo")
        history = self.db.query(RSHistory).filter_by(symbol="AAPL").one()
        points = decode_points(history.points_payload)
        self.assertLessEqual((points[-1][0] - points[0][0]).days, 545)
        self.assertGreaterEqual(history.row_count, 180)

    def test_excel_live_rows_and_workbook_for_both_markets(self):
        with patch.object(YahooProvider, "get_ohlcv", return_value=observations(years=6)):
            for symbol, exchange in [("AAPL", "US"), ("RELIANCE", "NSE")]:
                feed = market.get_excel_feed(symbol, exchange, 5000, self.db)
                self.assertGreater(len(feed["ohlcv"]), 1000)
                self.assertTrue(feed["history"]["requirement_met"])
                self.assertEqual(feed["history"]["stored_rows"], 0)
                self.assertIsNone(feed["fundamental_snapshot"]["market_cap"])
                result = market.export_excel_snapshot(symbol, exchange, 5000, self.db)
                book = load_workbook(BytesIO(result.body))
                self.assertEqual(book["Summary"]["B1"].value, symbol)
                self.assertEqual(book["Summary"]["B2"].value, exchange)
                self.assertGreater(book["OHLCV"].max_row, 1000)
                self.assertEqual(book["Fundamental_Ownership"]["B2"].value, "N/A")
                self.assertEqual(book["Ranking_Config"]["C2"].value, 30)
                self.assertEqual(book["Ranking_Config"]["B8"].data_type, "f")
                csv = market.get_excel_live_csv(symbol, exchange, 5000, self.db)
                self.assertIn(symbol.encode(), csv.body)

    def test_backtest_both_markets_real_history_contract(self):
        with patch.object(YahooProvider, "get_ohlcv", return_value=observations(years=20)):
            for symbol, exchange in [("AAPL", "US"), ("RELIANCE", "NSE")]:
                result = market.get_backtest(symbol, exchange, 20, self.db)
                self.assertGreater(result["actual_years"], 19)
                self.assertEqual(result["history_status"], "Complete")
                self.assertGreater(len(result["equity_curve"]), 500)
                self.assertEqual(len(result["strategies"]), 6)
                self.assertIsInstance(result["buy_hold_return_percent"], float)
                self.assertTrue(all(s["max_drawdown_percent"] <= 0 for s in result["strategies"]))

    def test_provider_failure_is_useful_not_server_exception(self):
        with patch.object(YahooProvider, "get_ohlcv", side_effect=RuntimeError("rate limited")), patch.object(YahooProvider, "get_chart_history", return_value=[]):
            with self.assertRaises(market.HTTPException) as caught:
                market.get_backtest("RELIANCE", "NSE", 20, self.db)
            self.assertEqual(caught.exception.status_code, 503)
            self.assertIn("NSE:RELIANCE", caught.exception.detail)

    def test_yahoo_cookie_failure_uses_same_provider_direct_history(self):
        expected = observations()
        with patch("app.services.providers.yahoo_provider.yf.Ticker") as ticker, patch.object(YahooProvider, "get_chart_history", return_value=expected) as direct:
            ticker.return_value.history.side_effect = RuntimeError("cookie error")
            self.assertEqual(YahooProvider().get_ohlcv("RELIANCE", "NSE", "2006-01-01"), expected)
            direct.assert_called_once_with("RELIANCE", "NSE", "2006-01-01", None)

    def test_chart_parser_never_fills_missing_ohlc(self):
        payload = {"timestamp": [1700000000, 1700100000], "meta": {"exchangeTimezoneName": "Asia/Kolkata"}, "indicators": {"quote": [{"open": [10, None], "high": [12, 12], "low": [9, 9], "close": [11, 11], "volume": [None, 20]}]}}
        parsed = YahooProvider.parse_chart_history(payload)
        self.assertEqual(len(parsed), 1)
        self.assertIsNone(parsed[0]["volume"])

    def test_fixed_5000_percentile_formula_ties(self):
        self.assertEqual(market._percentile_rank([1] * 5000, 1, 5000), 50)
        self.assertEqual(market._percentile_rank(list(range(5000)), 4999, 5000), 100)
        self.assertEqual(market._percentile_rank([1, 2, 3], 3, 5000), .04)

    def test_compact_history_survives_restart_and_no_ohlcv_writes(self):
        with patch.object(backfill, "SessionLocal", self.Session), patch.object(backfill.yf, "download", return_value=pd.DataFrame()), patch.object(YahooProvider, "get_ohlcv", return_value=observations()):
            result = backfill.backfill_market_batch("US", 1)
        self.assertEqual(result["succeeded"], 1)
        self.db.expire_all()
        self.assertEqual(self.db.query(OHLCV).count(), 0)
        item = self.db.query(RSHistory).one()
        self.assertGreater(len(decode_points(item.points_payload)), 300)
        backfill._last_results.clear()
        status = backfill.universe_backfill_status(self.db, "US")
        self.assertEqual(status["history_ready_symbols"], 1)
        self.assertEqual(status["last_batch"]["succeeded"], 1)
        self.assertFalse(status["complete"])

    def test_retries_persist_and_do_not_starve_next_stock(self):
        self.db.add(Company(symbol="AA", exchange="US", name="Test equity", is_active=1))
        self.db.commit()
        with patch.object(backfill, "SessionLocal", self.Session), patch.object(backfill.yf, "download", return_value=pd.DataFrame()), patch.object(YahooProvider, "get_ohlcv", side_effect=RuntimeError("limited")):
            backfill.backfill_market_batch("US", 1)
        backfill._symbol_retry_after.clear()
        self.db.expire_all()
        candidates = backfill._candidate_companies(self.db, "US", 1)
        self.assertEqual(candidates[0].symbol, "AAPL")

    def test_india_dual_listing_is_not_two_stocks(self):
        self.db.add(Company(symbol="RELIANCE", exchange="BSE", name="Reliance", isin="INE002A01018", is_active=1))
        self.db.commit()
        self.assertEqual(backfill.universe_backfill_status(self.db, "INDIA")["active_symbols"], 1)

    def test_live_rows_keep_symbol_identity(self):
        with patch.object(YahooProvider, "get_ohlcv", return_value=observations()):
            rows = market._live_provider_rows("AAPL", "US")
        self.assertEqual(rows[-1].symbol, "AAPL")
        self.assertEqual(rows[-1].exchange, "US")

    def test_rs_weights_only_change_score_not_raw_returns(self):
        rows = [SimpleNamespace(**r, symbol="AAPL", exchange="US") for r in observations()]
        bench = [(r.date, r.open * .8, r.close * .8) for r in rows]
        _, metrics = market._weighted_relative_return_from_points([(r.date, r.open, r.close) for r in rows], bench, {k: 1 for k in ("1w", "2w", "1m", "2m", "3m", "6m", "1y")})
        universe = {"US:AAPL": metrics, "US:OTHER": {k: {**v, "relative_return_percent": v["relative_return_percent"] + (1 if k == "1w" else -1)} for k, v in metrics.items()}}
        with patch.object(market, "_benchmark_close_points", return_value=(bench, "^GSPC", "S&P 500", "verified")), patch.object(market, "_rs_universe_metrics", return_value=(universe, {})):
            original = market._weighted_rs_against_benchmark(rows, "US", db=self.db)
            changed = market._weighted_rs_against_benchmark(rows, "US", {"1w": 100, "1m": 0, "3m": 0, "6m": 0, "1y": 0}, db=self.db)
        self.assertNotEqual(original[0], changed[0])
        self.assertEqual(original[1], changed[1])
        self.assertEqual(original[2]["1w"]["relative_return_percent"], changed[2]["1w"]["relative_return_percent"])
        self.assertEqual(original[2]["_scored_stocks_available"], 2)

    def test_compact_universe_is_scored_after_ohlcv_compaction(self):
        item = persist_history(self.db, "AAPL", "US", observations())
        points = decode_points(item.points_payload)
        metrics, _ = market._rs_universe_metrics(self.db, "US", points)
        self.assertIn("US:AAPL", metrics)
        self.assertEqual(set(metrics["US:AAPL"]), {"1w", "2w", "1m", "2m", "3m", "6m", "1y"})

    def test_full_5000_cohort_is_bounded_and_scores_available(self):
        item = persist_history(self.db, "AAPL", "US", observations())
        points = decode_points(item.points_payload)
        self.db.add_all([Company(symbol=f"ZZ{i:04d}", exchange="US", name=f"Test stock {i}", is_active=1) for i in range(5000)])
        self.db.add_all([RSHistory(symbol=f"ZZ{i:04d}", exchange="US", points_payload=item.points_payload, row_count=item.row_count, first_date=item.first_date, last_date=item.last_date, status="success") for i in range(5000)])
        self.db.commit()
        rows = [SimpleNamespace(**r, symbol="AAPL", exchange="US") for r in observations()]
        with patch.object(market, "_benchmark_close_points", return_value=(points, "^GSPC", "S&P 500", "verified")):
            result = market._weighted_rs_against_benchmark(rows, "US", db=self.db)
        self.assertEqual(result[0], 50)
        self.assertEqual(result[2]["_scored_stocks_available"], 5000)
        self.assertTrue(result[2]["_score_complete"])


    def test_bse_directory_excludes_funds_and_placeholder_isin(self):
        from app.services.providers.bse_provider import BSEProvider
        with patch("app.services.providers.bse_provider.requests.get") as get:
            get.return_value.json.return_value = {"status":"ok","data":[
                {"symbol":"TEST","name":"Test Industries","mic_code":"XBOM","type":"Common Stock","isin":"request_access_via_add_ons"},
                {"symbol":"FUND","name":"Test Mutual Fund","mic_code":"XBOM","type":"Common Stock"}]}
            rows = BSEProvider().get_companies()
        self.assertEqual(len(rows),1)
        self.assertIsNone(rows[0]["isin"])
        self.assertEqual(rows[0]["exchange"],"BSE")

    def test_india_missing_bse_isin_does_not_double_count_same_ticker(self):
        self.db.add(Company(symbol="RELIANCE",exchange="BSE",name="Reliance",is_active=1))
        self.db.commit()
        self.assertEqual(backfill.universe_backfill_status(self.db,"INDIA")["active_symbols"],1)
        self.assertEqual([c.exchange for c in backfill._candidate_companies(self.db,"INDIA",5)],["NSE"])


if __name__ == "__main__":
    unittest.main()
