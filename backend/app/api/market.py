from app.models import OHLCV

from fastapi import APIRouter, Depends, HTTPException
from fastapi import Query
from fastapi.responses import Response
from sqlalchemy.orm import Session, aliased
from sqlalchemy import func, or_, case

from app.database import SessionLocal
from app.services.providers.yahoo_provider import YahooProvider
from app.services.ohlcv_sync import sync_ohlcv
from app.services.fundamental_sync import sync_fundamental_data
from app.models import Fundamental, Ownership, Company
from app.services.providers.bse_provider import BSEProvider
from app.services.providers.india_shareholding_provider import IndiaShareholdingProvider
from app.services.providers.sec_provider import SECFundamentalsProvider
from app.services.scheduler import refresh_configured_market_data
from app.services.equity_filters import apply_eligible_equity_filter
from app.services.universe_data_backfill import universe_data_status

import os
import math
import requests
from io import BytesIO
from datetime import datetime, timedelta, date
import yfinance as yf

router = APIRouter(prefix="/market", tags=["market"])


def _json_safe(value):
    """Recursively replace non-finite numeric values with None for JSON responses."""
    if isinstance(value, dict):
        return {k: _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    # numpy/pandas floating scalars can reach API responses after resampling.
    try:
        if value.__class__.__module__.startswith(("numpy", "pandas")):
            numeric = float(value)
            return numeric if math.isfinite(numeric) else None
    except Exception:
        pass
    return value


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def _valid_trading_rows(rows, exchange: str):
    """Exclude invalid stock bars from calculations and API output.

    Normal US/NSE/BSE equity candles must be weekday-dated and have finite,
    strictly positive OHLC values.  Filtering here protects the UI and ranking
    calculations from legacy/provider rows such as a zero close.
    """
    rows = list(rows)
    if exchange.upper() not in {"US", "NSE", "BSE"}:
        return rows

    valid = []
    for row in rows:
        if row.date is None or row.date.weekday() >= 5:
            continue
        try:
            ohlc = [float(row.open), float(row.high), float(row.low), float(row.close)]
        except (TypeError, ValueError):
            continue
        if any((not math.isfinite(value)) or value <= 0 for value in ohlc):
            continue
        valid.append(row)
    return valid


def _ema(values, period):
    if len(values) < period:
        return None
    multiplier = 2 / (period + 1)
    value = sum(values[:period]) / period
    for price in values[period:]:
        value = ((price - value) * multiplier) + value
    return value


def _rma(values, period):
    """TradingView/Wilder moving average used by ATR and RSI-style smoothing."""
    if len(values) < period:
        return None
    value = sum(values[:period]) / period
    alpha = 1 / period
    for item in values[period:]:
        value = (alpha * item) + ((1 - alpha) * value)
    return value


def _rsi(values, period=14):
    if len(values) < period + 1:
        return None
    changes = [values[i] - values[i - 1] for i in range(1, len(values))]
    gains = [max(change, 0) for change in changes]
    losses = [max(-change, 0) for change in changes]

    def rma(items):
        result = sum(items[:period]) / period
        alpha = 1 / period
        for item in items[period:]:
            result = (alpha * item) + ((1 - alpha) * result)
        return result

    avg_gain = rma(gains)
    avg_loss = rma(losses)
    if avg_loss == 0:
        return 100.0
    rs = avg_gain / avg_loss
    return 100 - (100 / (1 + rs))


CLIENT_COMPOSITE_WEIGHTS = {
    "fundamental": 30.0,
    "technical": 25.0,
    "relative_strength": 25.0,
    "ownership": 15.0,
    "sector": 5.0,
}

CLIENT_SECTOR_WEIGHTS = {
    "eps_growth_rs": 30.0,
    "pat_growth_rs": 25.0,
    "sales_growth_rs": 20.0,
    "growth_acceleration_rs": 15.0,
    "growth_breadth": 5.0,
    "acceleration_breadth": 5.0,
}


# Compact filter-table defaults from the client's latest handwritten layout.
# Thresholds/weights that were not legible or not explicitly defined remain
# zero/disabled so the screener never invents a ranking rule.
CLIENT_TECHNICAL_FILTER_CONFIG = [
    {"name": "Upper BB - Lower BB", "compare": "<=", "value": 10.0, "weight": 10.0, "enabled": True},
    {"name": "5-day ATR% average vs 20-day ATR% average", "compare": "<", "value": "20-day ATR%", "weight": 10.0, "enabled": True},
    {"name": "10-day ATR% average vs 20-day ATR% average", "compare": "<", "value": "20-day ATR%", "weight": 5.0, "enabled": True},
    {"name": "RSI (14)", "compare": "bands", "value": "30 / 40 / 50", "weight": 5.0, "enabled": True},
    {"name": "10-day volume average vs 20-day volume average", "compare": "<", "value": "20-day volume", "weight": 10.0, "enabled": True},
    {"name": "20-day volume average vs 40-day volume average", "compare": "<", "value": "40-day volume", "weight": 5.0, "enabled": True},
    {"name": "Distance from 52-week high", "compare": "bands", "value": "10 / 17 / 20%", "weight": 10.0, "enabled": True},
    {"name": "20 EMA vs 50 EMA", "compare": ">", "value": "50 EMA", "weight": 8.0, "enabled": True},
    {"name": "50 EMA vs 150 EMA", "compare": ">", "value": "150 EMA", "weight": 4.0, "enabled": True},
]

CLIENT_FUNDAMENTAL_FILTER_CONFIG = [
    {"group": "EPS", "name": "Latest quarter EPS YoY", "compare": ">", "value": 30.0, "weight": 10.0, "enabled": True},
    {"group": "EPS", "name": "Quarterly EPS rising", "compare": ">", "value": "previous 2 quarters", "weight": 4.0, "enabled": True},
    {"group": "EPS", "name": "Quarterly EPS YoY trend rising", "compare": ">", "value": "previous 2 YoY values", "weight": 6.0, "enabled": True},
    {"group": "EPS", "name": "Latest year EPS YoY", "compare": ">", "value": 20.0, "weight": 6.0, "enabled": True},
    {"group": "EPS", "name": "Annual EPS trend rising", "compare": ">", "value": "previous 2 years", "weight": 4.0, "enabled": True},
    {"group": "PAT", "name": "Latest quarter PAT YoY", "compare": ">", "value": 30.0, "weight": 6.0, "enabled": True},
    {"group": "PAT", "name": "Quarterly PAT rising", "compare": ">", "value": "previous 2 quarters", "weight": 4.0, "enabled": True},
    {"group": "PAT", "name": "Quarterly PAT YoY trend rising", "compare": ">", "value": "previous 2 YoY values", "weight": 6.0, "enabled": True},
    {"group": "PAT", "name": "Latest year PAT YoY", "compare": ">", "value": 20.0, "weight": 5.0, "enabled": True},
    {"group": "PAT", "name": "Annual PAT trend rising", "compare": ">", "value": "previous 2 years", "weight": 4.0, "enabled": True},
    {"group": "Sales", "name": "Latest quarter Sales YoY", "compare": ">", "value": 30.0, "weight": 7.0, "enabled": True},
    {"group": "Sales", "name": "Quarterly Sales trend rising", "compare": ">", "value": "previous 2 quarters", "weight": 4.0, "enabled": True},
    {"group": "Sales", "name": "Latest year Sales YoY", "compare": ">", "value": 20.0, "weight": 5.0, "enabled": True},
    {"group": "Sales", "name": "Annual Sales trend rising", "compare": ">", "value": "previous 2 years", "weight": 4.0, "enabled": True},
    {"group": "NPM", "name": "Quarterly NPM growth YoY", "compare": ">", "value": 20.0, "weight": 4.0, "enabled": True},
    {"group": "NPM", "name": "Annual NPM rising", "compare": ">", "value": "previous year", "weight": 3.0, "enabled": True},
    {"group": "CFO", "name": "Operating cash flow YoY", "compare": ">", "value": 10.0, "weight": 4.0, "enabled": True},
    {"group": "CFO", "name": "Cash flow per share", "compare": ">", "value": 0.0, "weight": 0.0, "enabled": False},
    {"group": "Other", "name": "ROE", "compare": ">", "value": 20.0, "weight": 0.0, "enabled": True},
    {"group": "Other", "name": "ROCE", "compare": ">", "value": 30.0, "weight": 0.0, "enabled": True},
    {"group": "Other", "name": "Outstanding shares", "compare": "<", "value": 0.0, "weight": 0.0, "enabled": False},
    {"group": "Other", "name": "Float shares", "compare": "<", "value": 0.0, "weight": 0.0, "enabled": False},
]


def _roc(values, periods):
    if periods <= 0 or len(values) <= periods:
        return None
    base = values[-1 - periods]
    if base in (None, 0):
        return None
    return ((values[-1] / base) - 1) * 100


def _wilder_series(values, period):
    if len(values) < period:
        return []
    out = [None] * len(values)
    value = sum(values[:period]) / period
    out[period - 1] = value
    for i in range(period, len(values)):
        value = ((value * (period - 1)) + values[i]) / period
        out[i] = value
    return out


def _adx_di(highs, lows, closes, period=14):
    """Return latest ADX, +DI, -DI and DI spread using Wilder smoothing."""
    if len(closes) < (period * 2 + 1):
        return None, None, None, None
    trs, plus_dm, minus_dm = [], [], []
    for i in range(1, len(closes)):
        up = highs[i] - highs[i - 1]
        down = lows[i - 1] - lows[i]
        plus_dm.append(up if up > down and up > 0 else 0.0)
        minus_dm.append(down if down > up and down > 0 else 0.0)
        trs.append(max(
            highs[i] - lows[i],
            abs(highs[i] - closes[i - 1]),
            abs(lows[i] - closes[i - 1]),
        ))
    atr_s = _wilder_series(trs, period)
    plus_s = _wilder_series(plus_dm, period)
    minus_s = _wilder_series(minus_dm, period)
    dx = []
    di_pairs = []
    for tr, pdm, mdm in zip(atr_s, plus_s, minus_s):
        if tr in (None, 0) or pdm is None or mdm is None:
            dx.append(None)
            di_pairs.append((None, None))
            continue
        pdi = 100 * pdm / tr
        mdi = 100 * mdm / tr
        denom = pdi + mdi
        dx.append((100 * abs(pdi - mdi) / denom) if denom else 0.0)
        di_pairs.append((pdi, mdi))
    usable_dx = [v for v in dx if v is not None]
    if len(usable_dx) < period:
        return None, None, None, None
    adx_series = _wilder_series(usable_dx, period)
    adx = next((v for v in reversed(adx_series) if v is not None), None)
    plus_di, minus_di = next(((p, m) for p, m in reversed(di_pairs) if p is not None and m is not None), (None, None))
    spread = (plus_di - minus_di) if plus_di is not None and minus_di is not None else None
    return adx, plus_di, minus_di, spread


def _shift_months(value_date, months):
    import calendar
    month_index = value_date.year * 12 + (value_date.month - 1) - months
    year = month_index // 12
    month = (month_index % 12) + 1
    day = min(value_date.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


def _rs_anchor_date(end_date, label):
    if label == "1w":
        return end_date - timedelta(days=7)
    if label == "2w":
        return end_date - timedelta(days=14)
    if label == "1m":
        return _shift_months(end_date, 1)
    if label == "2m":
        return _shift_months(end_date, 2)
    if label == "3m":
        return _shift_months(end_date, 3)
    if label == "6m":
        return _shift_months(end_date, 6)
    if label == "1y":
        return _shift_months(end_date, 12)
    return end_date


def _point_date(point):
    return point[0]


def _point_open(point):
    if len(point) >= 3:
        return float(point[1]) if point[1] not in (None, 0) else None
    return None


def _point_close(point):
    value = point[2] if len(point) >= 3 else point[1]
    return float(value) if value not in (None, 0) else None


def _close_on_or_before(points, target_date):
    for point in reversed(points):
        trade_date = _point_date(point)
        close = _point_close(point)
        if trade_date <= target_date and close is not None:
            return trade_date, close
    return None


def _close_on_or_after(points, target_date):
    for point in points:
        trade_date = _point_date(point)
        close = _point_close(point)
        if trade_date >= target_date and close is not None:
            return trade_date, close
    return None


def _period_start_date(end_date, label):
    """Start of the current TradingView-style candle for the requested period.

    The client compares the percent shown on TradingView after selecting 1W/1M/etc.
    That percentage is the current candle's open-to-current-close change, not a
    trailing close-to-close return.
    """
    if label == "1d":
        return end_date
    if label == "1w":
        return end_date - timedelta(days=end_date.weekday())
    if label == "2w":
        monday = end_date - timedelta(days=end_date.weekday())
        return monday - timedelta(days=7)
    if label == "3w":
        monday = end_date - timedelta(days=end_date.weekday())
        return monday - timedelta(days=14)
    if label == "1m":
        return end_date.replace(day=1)
    if label == "2m":
        month_index = end_date.month - 1
        start_month = (month_index // 2) * 2 + 1
        return end_date.replace(month=start_month, day=1)
    if label == "3m":
        start_month = ((end_date.month - 1) // 3) * 3 + 1
        return end_date.replace(month=start_month, day=1)
    if label == "6m":
        start_month = 1 if end_date.month <= 6 else 7
        return end_date.replace(month=start_month, day=1)
    if label == "1y":
        return end_date.replace(month=1, day=1)
    return end_date


def _open_on_or_after(points, target_date):
    for point in points:
        trade_date = _point_date(point)
        if trade_date < target_date:
            continue
        open_value = _point_open(point)
        if open_value is not None:
            return trade_date, open_value
        close_value = _point_close(point)
        if close_value is not None:
            return trade_date, close_value
    return None


def _period_returns_from_points(points, labels=("1d", "1w", "2w", "3w", "1m", "2m", "3m", "6m", "1y")):
    """Return TradingView-style current-candle returns for comparison tables."""
    if not points:
        return {}
    points = sorted(points, key=_point_date)
    end_date = _point_date(points[-1])
    end = _close_on_or_before(points, end_date)
    if not end:
        return {}
    result = {}
    for label in labels:
        start = _open_on_or_after(points, _period_start_date(end_date, label))
        if not start or start[1] in (None, 0):
            result[label] = None
            continue
        result[label] = round(((end[1] / start[1]) - 1) * 100, 2)
    return result


def _yf_period_returns(ticker_symbol):
    try:
        hist = yf.Ticker(ticker_symbol).history(period="2y", interval="1d", auto_adjust=False)
        points = [
            (idx.date(), float(row["Open"]), float(row["Close"]))
            for idx, row in hist.iterrows()
            if row.get("Open") is not None and row.get("Close") is not None
            and math.isfinite(float(row["Open"])) and math.isfinite(float(row["Close"]))
        ]
        return _period_returns_from_points(points)
    except Exception:
        return {}


def _peer_group_period_returns(db, exchange, field_name, field_value, exclude_symbol=None, minimum_peers=5):
    """Average raw returns for real stored peers only.

    The selected stock is excluded. If the database does not contain enough
    peer histories, return no values so the UI shows N/A instead of repeating
    the selected stock's own return as an apparent sector/industry return.
    """
    if db is None or not field_value or field_name not in {"sector", "industry"}:
        return {}
    field = getattr(Company, field_name)
    peer_symbols = [
        row.symbol for row in db.query(Company.symbol)
        .filter(Company.exchange == exchange.upper(), field == field_value, Company.is_active == 1)
        .all()
    ]
    if exclude_symbol:
        peer_symbols = [s for s in peer_symbols if s != exclude_symbol.upper()]
    if len(peer_symbols) < minimum_peers:
        return {}

    cutoff = date.today() - timedelta(days=430)
    rows = (
        db.query(OHLCV)
        .filter(OHLCV.exchange == exchange.upper(), OHLCV.symbol.in_(peer_symbols), OHLCV.date >= cutoff)
        .order_by(OHLCV.symbol.asc(), OHLCV.date.asc())
        .all()
    )
    grouped = {}
    for row in _valid_trading_rows(rows, exchange):
        if row.close is None:
            continue
        grouped.setdefault(row.symbol, []).append(
            (row.date, float(row.open) if row.open is not None else float(row.close), float(row.close))
        )

    per_symbol = [
        _period_returns_from_points(points)
        for points in grouped.values()
        if len(points) >= 2
    ]
    if len(per_symbol) < minimum_peers:
        return {}

    result = {}
    for label in ("1d", "1w", "2w", "3w", "1m", "2m", "3m", "6m", "1y"):
        vals = [item.get(label) for item in per_symbol if item.get(label) is not None]
        result[label] = round(sum(vals) / len(vals), 2) if vals else None
    return result


def _nse_delivery_summary(symbol):
    """Best-effort NSE delivery percentage; never substitutes ordinary volume for delivery data."""
    try:
        session = requests.Session()
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/123 Safari/537.36",
            "Accept-Language": "en-US,en;q=0.9",
            "Referer": "https://www.nseindia.com/",
        }
        session.get("https://www.nseindia.com/", headers=headers, timeout=8)
        end = date.today()
        start = end - timedelta(days=45)
        params = {
            "symbol": symbol.upper(),
            "series": '["EQ"]',
            "from": start.strftime("%d-%m-%Y"),
            "to": end.strftime("%d-%m-%Y"),
        }
        response = session.get("https://www.nseindia.com/api/historical/cm/equity", params=params, headers=headers, timeout=10)
        response.raise_for_status()
        payload = response.json()
        data = payload.get("data") or []
        parsed = []
        for item in data:
            pct = item.get("CH_DELIV_PER")
            qty = item.get("CH_DELIV_QTY")
            total = item.get("CH_TOT_TRADED_QTY")
            dt = item.get("CH_TIMESTAMP") or item.get("mTIMESTAMP")
            try:
                pct = float(pct) if pct not in (None, "", "-") else None
                qty = float(qty) if qty not in (None, "", "-") else None
                total = float(total) if total not in (None, "", "-") else None
            except Exception:
                continue
            parsed.append({"date": str(dt or ""), "percent": pct, "delivered": qty, "total": total})
        if not parsed:
            return {"available": False, "source": "NSE", "note": "NSE delivery data unavailable"}

        def aggregate(items):
            delivered = sum(x["delivered"] for x in items if x["delivered"] is not None)
            total = sum(x["total"] for x in items if x["total"] is not None)
            if delivered and total:
                pct = delivered / total * 100
            else:
                vals = [x["percent"] for x in items if x["percent"] is not None]
                pct = sum(vals) / len(vals) if vals else None
            return {
                "percent": round(pct, 2) if pct is not None else None,
                "delivered_quantity": round(delivered, 0) if delivered else None,
                "traded_quantity": round(total, 0) if total else None,
            }

        # NSE generally returns newest first.
        return {
            "available": True,
            "source": "NSE",
            "day": aggregate(data and parsed[:1] or []),
            "weekly": aggregate(parsed[:5]),
            "monthly": aggregate(parsed[:20]),
        }
    except Exception:
        # NSE may return an HTML/block page instead of JSON from cloud-hosted
        # servers. Do not surface parser/internal exception text to the user.
        return {
            "available": False,
            "source": "NSE",
            "note": "Delivery percentage is currently unavailable from the live NSE delivery-data endpoint.",
        }


def _weighted_relative_return_from_points(stock_points, benchmark_points, weights):
    """Client method using TradingView-style period candle returns.

    Period return = (current close / current period candle open - 1) * 100.
    Period Relative Return = Stock Return % - Benchmark Return %.
    """
    if not stock_points or not benchmark_points:
        return None, {}

    end_date = min(_point_date(stock_points[-1]), _point_date(benchmark_points[-1]))
    stock_end = _close_on_or_before(stock_points, end_date)
    bench_end = _close_on_or_before(benchmark_points, end_date)
    if not stock_end or not bench_end:
        return None, {}

    metrics = {}
    weighted_sum = 0.0
    available_weight = 0.0
    for label in ("1w", "2w", "1m", "2m", "3m", "6m", "1y"):
        weight = max(0.0, float(weights.get(label, 0) or 0))
        period_start = _period_start_date(end_date, label)
        stock_old = _open_on_or_after(stock_points, period_start)
        bench_old = _open_on_or_after(benchmark_points, period_start)
        if not stock_old or not bench_old or stock_old[1] == 0 or bench_old[1] == 0:
            continue

        stock_return = ((stock_end[1] / stock_old[1]) - 1) * 100
        benchmark_return = ((bench_end[1] / bench_old[1]) - 1) * 100
        relative_return = stock_return - benchmark_return
        metrics[label] = {
            "weight_percent": weight,
            "stock_return_percent": round(stock_return, 2),
            "benchmark_return_percent": round(benchmark_return, 2),
            "relative_return_percent": round(relative_return, 2),
            "start_date": stock_old[0].isoformat(),
            "benchmark_start_date": bench_old[0].isoformat(),
            "period_start_date": period_start.isoformat(),
            "end_date": stock_end[0].isoformat(),
            "return_method": "period_open_to_current_close",
        }
        if weight > 0:
            weighted_sum += relative_return * weight
            available_weight += weight

    if available_weight <= 0:
        return None, metrics
    return weighted_sum / available_weight, metrics

RS_PERCENTILE_TARGETS = {
    "US": 6000,
    "INDIA": 5500,
}


def _rs_market_group(exchange: str) -> str:
    return "US" if exchange.upper() == "US" else "INDIA"


def _rs_market_exchanges(exchange: str):
    return ("US",) if exchange.upper() == "US" else ("NSE", "BSE")


def _rs_percentile_total(exchange: str) -> int:
    return RS_PERCENTILE_TARGETS[_rs_market_group(exchange)]

def _percentile_rank(values, target_value, total_count=None):
    """Client percentile: (lower + 0.5 * equal) * 100 / total.

    Stock RS uses the client-defined fixed market denominator: 6,000 for the
    US universe and 5,500 for the combined Indian (NSE/BSE) universe. Other
    percentile uses (for example sector-to-sector comparisons) can leave
    ``total_count`` unset and use the actual sample size.
    """
    values = [float(v) for v in values if v is not None and math.isfinite(float(v))]
    if target_value is None or not values:
        return None
    tolerance = 1e-9
    lower = sum(1 for value in values if value < target_value - tolerance)
    equal = sum(1 for value in values if abs(value - target_value) <= tolerance)
    denominator = int(total_count) if total_count is not None else len(values)
    if denominator <= 0:
        return None
    percentile = ((lower + 0.5 * equal) * 100.0) / denominator
    return round(max(0.0, min(100.0, percentile)), 2)


def _rs_universe_metrics(db: Session, exchange: str, benchmark_points):
    """Return period relative returns for the selected market universe.

    US scores use US listings only. Indian scores use stored NSE + BSE listings
    together, matching the client's requirement for one Indian RS universe.
    Exchange is included in the internal key so equal ticker strings on two
    exchanges cannot overwrite one another.
    """
    if db is None:
        return {}, {}

    market_exchanges = _rs_market_exchanges(exchange)
    cutoff = date.today() - timedelta(days=430)
    rows = (
        db.query(OHLCV)
        .join(
            Company,
            (Company.symbol == OHLCV.symbol) & (Company.exchange == OHLCV.exchange),
        )
        .filter(
            OHLCV.exchange.in_(market_exchanges),
            OHLCV.date >= cutoff,
            Company.is_active == 1,
        )
        .order_by(OHLCV.exchange.asc(), OHLCV.symbol.asc(), OHLCV.date.asc())
        .all()
    )
    grouped = {}
    for row in rows:
        if row.date is None or row.date.weekday() >= 5 or row.close is None:
            continue
        key = f"{row.exchange.upper()}:{row.symbol.upper()}"
        grouped.setdefault(key, []).append(
            (row.date, float(row.open) if row.open is not None else float(row.close), float(row.close))
        )

    # The relative-return metrics themselves do not depend on scoring weights.
    metric_weights = {key: 1.0 for key in ("1w", "2w", "1m", "2m", "3m", "6m", "1y")}
    symbol_metrics = {}
    for key, points in grouped.items():
        _, metrics = _weighted_relative_return_from_points(points, benchmark_points, metric_weights)
        if metrics:
            symbol_metrics[key] = metrics

    sector_by_symbol = {
        f"{row.exchange.upper()}:{row.symbol.upper()}": row.sector
        for row in db.query(Company)
        .filter(Company.exchange.in_(market_exchanges), Company.is_active == 1)
        .all()
        if row.sector
    }
    return symbol_metrics, sector_by_symbol


def _weighted_rs_against_benchmark(daily_rows, exchange: str, period_weights=None, db: Session = None):
    """Client RS method: period relative returns -> period percentiles -> weighted RS score."""
    benchmark_symbol = "^GSPC" if exchange.upper() == "US" else "^CRSLDX"
    benchmark_name = "S&P 500" if exchange.upper() == "US" else "NIFTY 500"

    # Latest client handwritten RS reference:
    # 1W*0.30 + 1M*0.25 + 3M*0.20 + 6M*0.15 + 12M*0.10
    # 2W/2M and Sector RS remain optional custom components with zero default weight.
    defaults = {"1w": 30.0, "2w": 0.0, "1m": 25.0, "2m": 0.0, "3m": 20.0, "6m": 15.0, "1y": 10.0, "sector": 0.0}
    weights = defaults.copy()
    if period_weights:
        for key, value in period_weights.items():
            if key in weights:
                try:
                    weights[key] = max(0.0, float(value))
                except Exception:
                    pass

    try:
        hist = yf.Ticker(benchmark_symbol).history(period="5y", interval="1d", auto_adjust=False)
        benchmark_points = [
            (idx.date(), float(row["Open"]), float(row["Close"]))
            for idx, row in hist.iterrows()
            if row.get("Open") is not None and row.get("Close") is not None
            and math.isfinite(float(row["Open"])) and math.isfinite(float(row["Close"]))
        ]
        stock_points = [
            (r.date, float(r.open) if r.open is not None else float(r.close), float(r.close))
            for r in daily_rows
            if r.date is not None and r.close is not None and math.isfinite(float(r.close))
        ]
        stock_points.sort(key=lambda item: item[0])
        benchmark_points.sort(key=lambda item: item[0])

        benchmark_by_date = {d: c for d, _o, c in benchmark_points}
        aligned = [(d, c, benchmark_by_date.get(d)) for d, _o, c in stock_points if benchmark_by_date.get(d) is not None]
        rs_chart = []
        base_ratio = None
        for trade_date, stock_close, bench_close in aligned:
            if not stock_close or not bench_close:
                continue
            ratio = stock_close / bench_close
            if base_ratio is None:
                base_ratio = ratio
            if base_ratio:
                rs_chart.append({"date": trade_date.isoformat(), "rs": round((ratio / base_ratio) * 100, 4)})

        metric_weights = {key: 1.0 for key in ("1w", "2w", "1m", "2m", "3m", "6m", "1y")}
        _, metrics = _weighted_relative_return_from_points(stock_points, benchmark_points, metric_weights)
        if not metrics:
            return None, None, metrics, benchmark_name, rs_chart

        target_size = _rs_percentile_total(exchange)
        universe_metrics, sector_by_symbol = _rs_universe_metrics(db, exchange, benchmark_points)
        for label in ("1w", "2w", "1m", "2m", "3m", "6m", "1y"):
            target_rr = (metrics.get(label) or {}).get("relative_return_percent")
            universe_values = [
                m[label]["relative_return_percent"]
                for m in universe_metrics.values()
                if label in m and m[label].get("relative_return_percent") is not None
            ]
            pct = _percentile_rank(universe_values, target_rr, target_size)
            if label in metrics:
                metrics[label]["percentile"] = pct
                metrics[label]["universe_size"] = target_size
                metrics[label]["scored_stocks_available"] = len(universe_values)
                metrics[label]["score_weight_percent"] = weights.get(label, 0.0)

        # First calculate each stock's period-percentile composite (without sector).
        period_labels = ("1w", "2w", "1m", "2m", "3m", "6m", "1y")
        period_universe_values = {
            label: [
                m[label]["relative_return_percent"]
                for m in universe_metrics.values()
                if label in m and m[label].get("relative_return_percent") is not None
            ]
            for label in period_labels
        }

        stock_period_scores = {}
        for symbol, sym_metrics in universe_metrics.items():
            score_sum = 0.0
            score_weight = 0.0
            for label in period_labels:
                weight = weights.get(label, 0.0)
                if weight <= 0 or label not in sym_metrics:
                    continue
                rr = sym_metrics[label].get("relative_return_percent")
                pct = _percentile_rank(period_universe_values[label], rr, target_size)
                if pct is None:
                    continue
                score_sum += pct * weight
                score_weight += weight
            if score_weight > 0:
                stock_period_scores[symbol] = score_sum / score_weight

        # Sector RS Score = percentile rank of the target sector's average member RS
        # versus the average RS of all other represented sectors.
        sector_groups = {}
        for symbol, score in stock_period_scores.items():
            sector = sector_by_symbol.get(symbol)
            if sector:
                sector_groups.setdefault(sector, []).append(score)
        sector_scores = {sector: sum(vals) / len(vals) for sector, vals in sector_groups.items() if vals}

        target_symbol = daily_rows[-1].symbol if daily_rows else None
        target_key = f"{exchange.upper()}:{target_symbol.upper()}" if target_symbol else None
        target_sector = sector_by_symbol.get(target_key) if target_key else None
        target_sector_raw = sector_scores.get(target_sector) if target_sector else None
        sector_percentile = _percentile_rank(list(sector_scores.values()), target_sector_raw)
        metrics["sector"] = {
            "sector": target_sector,
            "score": round(target_sector_raw, 2) if target_sector_raw is not None else None,
            "percentile": sector_percentile,
            "universe_size": len(sector_scores),
            "score_weight_percent": weights.get("sector", 0.0),
        }

        weighted_score = 0.0
        available_weight = 0.0
        for label in period_labels:
            weight = weights.get(label, 0.0)
            pct = (metrics.get(label) or {}).get("percentile")
            if weight > 0 and pct is not None:
                weighted_score += pct * weight
                available_weight += weight
        # Sector RS is an optional scoring component. Its default weight is 0,
        # so the standard period-only formula is unchanged unless the user enables it.
        sector_weight = weights.get("sector", 0.0)
        if sector_weight > 0 and sector_percentile is not None:
            weighted_score += sector_percentile * sector_weight
            available_weight += sector_weight

        if available_weight <= 0:
            return None, None, metrics, benchmark_name, rs_chart

        rating = round(weighted_score / available_weight, 2)

        # IMPORTANT: raw/composite relative return is intentionally independent
        # of the editable score weights.  The client specifically requested that
        # changing score weightage must change only Final RS Score.  Keep this
        # compatibility field on the fixed reference mix, while the final score
        # above uses the editable percentile weights.
        fixed_relative_weights = {
            "1w": 30.0, "2w": 0.0, "1m": 25.0, "2m": 0.0,
            "3m": 20.0, "6m": 15.0, "1y": 10.0,
        }
        weighted_relative, _ = _weighted_relative_return_from_points(
            stock_points, benchmark_points, fixed_relative_weights
        )
        metrics["_universe_size"] = target_size
        required_default_periods = ("1w", "1m", "3m", "6m", "1y")
        metrics["_scored_stocks_available"] = sum(
            1
            for symbol_metrics in universe_metrics.values()
            if all(
                label in symbol_metrics
                and symbol_metrics[label].get("relative_return_percent") is not None
                for label in required_default_periods
            )
        )
        metrics["_score_weight_total"] = available_weight
        metrics["_relative_return_weight_independent"] = True
        metrics["_fixed_relative_return_reference_weights"] = fixed_relative_weights
        return rating, weighted_relative, metrics, benchmark_name, rs_chart
    except Exception:
        return None, None, {}, benchmark_name, []


def _normalize_score_weights(weights=None):
    defaults = dict(CLIENT_COMPOSITE_WEIGHTS)
    if not weights:
        return defaults

    merged = {**defaults, **{k: max(0.0, float(v)) for k, v in weights.items() if k in defaults}}
    total = sum(merged.values())
    if total <= 0:
        return defaults
    return {k: (v / total) * 100.0 for k, v in merged.items()}


def _score_symbol(db: Session, symbol: str, exchange: str, weights=None, include_components=False, subweights=None, rs_period_weights=None):
    rows = (
        db.query(OHLCV)
        .filter(OHLCV.symbol == symbol, OHLCV.exchange == exchange)
        .order_by(OHLCV.date.asc())
        .all()
    )
    rows = _valid_trading_rows(rows, exchange)
    fundamental = (
        db.query(Fundamental)
        .filter(Fundamental.symbol == symbol, Fundamental.exchange == exchange)
        .first()
    )
    ownership = (
        db.query(Ownership)
        .filter(Ownership.symbol == symbol, Ownership.exchange == exchange)
        .first()
    )

    components = {
        "technical": None,
        "fundamental": None,
        "relative_strength": None,
        "ownership": None,
        # Sector growth score is a separate Milestone-2 component. It remains
        # unavailable until the peer fundamental-growth universe has enough
        # real observations; it is never replaced with a price-RS proxy.
        "sector": None,
    }

    closes = [float(row.close) for row in rows if row.close is not None]
    highs = [float(row.high) for row in rows if row.high is not None]
    lows = [float(row.low) for row in rows if row.low is not None]
    volumes = [float(row.volume or 0) for row in rows]

    if closes:
        latest = closes[-1]
        configured_subweights = subweights or {}
        technical_weights = configured_subweights.get("technical", {
            "ema20": 20, "ema50": 20, "ema150": 20, "ema200": 20, "rsi": 20,
        })
        technical_metrics = {}
        for period, key in [(20, "ema20"), (50, "ema50"), (150, "ema150"), (200, "ema200")]:
            ema = _ema(closes, period)
            if ema is not None:
                technical_metrics[key] = 100.0 if latest > ema else 0.0

        rsi = _rsi(closes, 14)
        if rsi is not None:
            technical_metrics["rsi"] = 100.0 if 50 <= rsi <= 70 else 50.0 if (40 <= rsi < 50 or 70 < rsi <= 80) else 0.0

        tw_sum = sum(max(0.0, float(technical_weights.get(k, 0))) for k in technical_metrics)
        if tw_sum > 0:
            components["technical"] = round(sum(technical_metrics[k] * max(0.0, float(technical_weights.get(k, 0))) for k in technical_metrics) / tw_sum, 2)

        # Keep dashboard RS aligned with the customizable Technical Summary model.
        rs_component, _, _, _, _ = _weighted_rs_against_benchmark(rows, exchange, rs_period_weights, db=db)
        if rs_component is not None:
            components["relative_strength"] = rs_component


    if fundamental:
        fundamental_weights = (subweights or {}).get("fundamental", {
            "eps": 20, "net_income": 20, "profit_margin": 20, "roe": 20, "roa": 20,
        })
        raw_checks = {
            "eps": (fundamental.trailing_eps, lambda x: x > 0),
            "net_income": (fundamental.net_income, lambda x: x > 0),
            "profit_margin": (fundamental.profit_margin, lambda x: x > 0),
            "roe": (fundamental.return_on_equity, lambda x: x >= 0.15),
            "roa": (fundamental.return_on_assets, lambda x: x >= 0.05),
        }
        fundamental_metrics = {}
        for key, (value, check) in raw_checks.items():
            if value is None:
                continue
            numeric = float(value)
            fundamental_metrics[key] = 100.0 if check(numeric) else (50.0 if numeric > 0 else 0.0)
        fw_sum = sum(max(0.0, float(fundamental_weights.get(k, 0))) for k in fundamental_metrics)
        if fw_sum > 0:
            components["fundamental"] = round(sum(fundamental_metrics[k] * max(0.0, float(fundamental_weights.get(k, 0))) for k in fundamental_metrics) / fw_sum, 2)

    if ownership:
        ownership_weights = (subweights or {}).get("ownership", {"institution": 70, "insider": 30})
        ownership_metrics = {}
        if ownership.institution_percent is not None:
            pct = float(ownership.institution_percent)
            if pct <= 1:
                pct *= 100
            ownership_metrics["institution"] = max(0, min(100, pct))
        if ownership.insider_percent is not None:
            pct = float(ownership.insider_percent)
            if pct <= 1:
                pct *= 100
            ownership_metrics["insider"] = max(0, min(100, pct * 5))
        ow_sum = sum(max(0.0, float(ownership_weights.get(k, 0))) for k in ownership_metrics)
        if ow_sum > 0:
            components["ownership"] = round(sum(ownership_metrics[k] * max(0.0, float(ownership_weights.get(k, 0))) for k in ownership_metrics) / ow_sum, 2)

    normalized_weights = _normalize_score_weights(weights)
    weighted_points = 0.0
    available_weight = 0.0
    for key, weight in normalized_weights.items():
        value = components.get(key)
        if value is None:
            continue
        weighted_points += value * weight
        available_weight += weight

    if available_weight <= 0:
        return (None, 0, components, normalized_weights) if include_components else (None, 0)

    score = round(weighted_points / available_weight)
    coverage = round(available_weight)
    result = (max(0, min(100, score)), max(0, min(100, coverage)))
    if include_components:
        return result[0], result[1], components, normalized_weights
    return result


def _rank_within(db: Session, company: Company, field: str, minimum_peers: int = 5):
    value = getattr(company, field)
    if not value:
        return {
            "available": False,
            "rank": None,
            "total": 0,
            "group": None,
            "note": "Classification data is unavailable for this stock."
        }

    peers = (
        db.query(Company)
        .filter(
            Company.exchange == company.exchange,
            getattr(Company, field) == value,
            Company.is_active == 1
        )
        .limit(100)
        .all()
    )

    scored = []
    for peer in peers:
        score, _ = _score_symbol(db, peer.symbol, peer.exchange)
        if score is not None:
            scored.append((peer.symbol, score))

    if len(scored) < minimum_peers:
        return {
            "available": False,
            "rank": None,
            "total": len(scored),
            "group": value,
            "note": f"Insufficient peer data ({len(scored)}/{minimum_peers} minimum)."
        }

    scored.sort(key=lambda item: item[1], reverse=True)
    for index, (peer_symbol, _) in enumerate(scored, start=1):
        if peer_symbol == company.symbol:
            return {
                "available": True,
                "rank": index,
                "total": len(scored),
                "group": value,
                "note": None
            }

    return {
        "available": False,
        "rank": None,
        "total": len(scored),
        "group": value,
        "note": "The selected stock does not yet have enough comparable scored peer data."
    }


def _stored_rs_rating(db: Session, symbol: str, exchange: str, minimum_universe: int = 20):
    cutoff = date.today() - timedelta(days=220)
    rows = (
        db.query(OHLCV)
        .filter(OHLCV.exchange == exchange, OHLCV.date >= cutoff)
        .order_by(OHLCV.symbol.asc(), OHLCV.date.asc())
        .limit(50000)
        .all()
    )
    rows = _valid_trading_rows(rows, exchange)

    grouped = {}
    for row in rows:
        if row.close is None:
            continue
        grouped.setdefault(row.symbol, []).append(float(row.close))

    returns = {}
    for peer_symbol, closes in grouped.items():
        if len(closes) >= 40 and closes[0] != 0:
            returns[peer_symbol] = (closes[-1] / closes[0]) - 1

    if symbol not in returns or len(returns) < minimum_universe:
        return None, len(returns)

    target = returns[symbol]
    below_or_equal = sum(1 for value in returns.values() if value <= target)
    percentile = round((below_or_equal / len(returns)) * 99)
    return max(1, min(99, percentile)), len(returns)


def _cached_symbol_summary(db: Session, symbol: str, exchange: str):
    rows = (
        db.query(OHLCV)
        .filter(
            OHLCV.symbol == symbol.upper(),
            OHLCV.exchange == exchange.upper(),
        )
        .order_by(OHLCV.date.desc())
        .limit(5000)
        .all()
    )
    valid = _valid_trading_rows(rows, exchange)
    if not valid:
        return None
    latest = valid[0]
    return {
        "cached_records": len(valid),
        "latest_date": latest.date.isoformat() if latest.date else None,
        "latest_close": float(latest.close) if latest.close is not None else None,
    }


@router.post("/refresh/{symbol}")
def refresh_symbol(
    symbol: str,
    exchange: str = "US",
    db: Session = Depends(get_db)
):
    symbol = symbol.upper().strip()
    exchange = exchange.upper().strip()

    def cached_response(reason: str):
        cached = _cached_symbol_summary(db, symbol, exchange)
        if not cached:
            return None
        return {
            "status": "cached",
            "provider_refresh_ok": False,
            "symbol": symbol,
            "exchange": exchange,
            "records_received": 0,
            "added": 0,
            "updated": 0,
            "removed_invalid_dates": 0,
            "warning": reason,
            **cached,
        }

    try:
        if exchange == "BSE":
            try:
                rows = BSEProvider().get_ohlcv(symbol)
            except Exception as primary_exc:
                # BSE has a second configured provider path through Yahoo.
                # If the paid/feed provider is temporarily unavailable, try the
                # alternate real-data source before falling back to stored bars.
                try:
                    rows = YahooProvider().get_ohlcv(
                        symbol=symbol,
                        exchange=exchange,
                        start_date="2000-01-01",
                    )
                except Exception:
                    cached = cached_response(f"Live BSE refresh unavailable: {primary_exc}")
                    if cached:
                        return cached
                    raise
        else:
            rows = YahooProvider().get_ohlcv(
                symbol=symbol,
                exchange=exchange,
                start_date="2000-01-01"
            )

        if not rows:
            cached = cached_response("The live provider returned no new rows; verified stored market data is being shown.")
            if cached:
                return cached
            raise HTTPException(status_code=404, detail="No market data returned and no stored market data is available")

        result = sync_ohlcv(
            db=db,
            symbol=symbol,
            exchange=exchange,
            rows=rows
        )

        cached = _cached_symbol_summary(db, symbol, exchange) or {}
        return {
            "status": "success",
            "provider_refresh_ok": True,
            "symbol": symbol,
            "exchange": exchange,
            "records_received": len(rows),
            **result,
            **cached,
        }

    except HTTPException:
        raise

    except Exception as exc:
        cached = cached_response(f"Live provider refresh failed temporarily: {exc}")
        if cached:
            return cached
        raise HTTPException(
            status_code=503,
            detail=f"Market data refresh failed and no verified stored data is available: {str(exc)}"
        )

@router.get("/history/{symbol}")
def get_history(
    symbol: str,
    exchange: str = "US",
    limit: int = 260,
    db: Session = Depends(get_db)
):
    rows = (
        db.query(OHLCV)
        .filter(
            OHLCV.symbol == symbol.upper(),
            OHLCV.exchange == exchange.upper()
        )
        .order_by(OHLCV.date.desc())
        .limit(limit)
        .all()
    )
    rows = _valid_trading_rows(rows, exchange)

    return [
        {
            "date": row.date,
            "open": row.open,
            "high": row.high,
            "low": row.low,
            "close": row.close,
            "volume": row.volume
        }
        for row in rows
    ]

@router.get("/dashboard/{symbol}")
def get_dashboard_summary(
    symbol: str,
    exchange: str = "US",
    technical_weight: float = 25,
    fundamental_weight: float = 30,
    relative_strength_weight: float = 25,
    ownership_weight: float = 15,
    sector_weight: float = 5,
    technical_ema20_weight: float = 20,
    technical_ema50_weight: float = 20,
    technical_ema150_weight: float = 20,
    technical_ema200_weight: float = 20,
    technical_rsi_weight: float = 20,
    fundamental_eps_weight: float = 20,
    fundamental_net_income_weight: float = 20,
    fundamental_profit_margin_weight: float = 20,
    fundamental_roe_weight: float = 20,
    fundamental_roa_weight: float = 20,
    ownership_institution_weight: float = 70,
    ownership_insider_weight: float = 30,
    rs_1w_weight: float = 30,
    rs_2w_weight: float = 0,
    rs_1m_weight: float = 25,
    rs_2m_weight: float = 0,
    rs_3m_weight: float = 20,
    rs_6m_weight: float = 15,
    rs_1y_weight: float = 10,
    rs_sector_weight: float = 0,
    db: Session = Depends(get_db)
):
    symbol = symbol.upper()
    exchange = exchange.upper()

    company = (
        db.query(Company)
        .filter(Company.symbol == symbol, Company.exchange == exchange)
        .first()
    )

    requested_weights = {
        "technical": technical_weight,
        "fundamental": fundamental_weight,
        "relative_strength": relative_strength_weight,
        "ownership": ownership_weight,
        "sector": sector_weight,
    }
    ranking_subweights = {
        "technical": {"ema20": technical_ema20_weight, "ema50": technical_ema50_weight, "ema150": technical_ema150_weight, "ema200": technical_ema200_weight, "rsi": technical_rsi_weight},
        "fundamental": {"eps": fundamental_eps_weight, "net_income": fundamental_net_income_weight, "profit_margin": fundamental_profit_margin_weight, "roe": fundamental_roe_weight, "roa": fundamental_roa_weight},
        "ownership": {"institution": ownership_institution_weight, "insider": ownership_insider_weight},
    }
    rs_period_weights = {"1w": rs_1w_weight, "2w": rs_2w_weight, "1m": rs_1m_weight, "2m": rs_2m_weight, "3m": rs_3m_weight, "6m": rs_6m_weight, "1y": rs_1y_weight, "sector": rs_sector_weight}
    score, coverage, components, normalized_weights = _score_symbol(
        db, symbol, exchange, weights=requested_weights, include_components=True,
        subweights=ranking_subweights, rs_period_weights=rs_period_weights
    )
    if score is None:
        raise HTTPException(status_code=404, detail="Not enough data to calculate dashboard score")

    # Do not present a seemingly complete Indian-market ranking when weighted
    # fundamental/ownership categories are unavailable. This directly exposes
    # the data-coverage limitation instead of silently re-normalizing it away.
    missing_required = []
    if exchange in {"NSE", "BSE"}:
        if normalized_weights.get("fundamental", 0) > 0 and components.get("fundamental") is None:
            missing_required.append("fundamental")
        if normalized_weights.get("ownership", 0) > 0 and components.get("ownership") is None:
            missing_required.append("ownership")

    if missing_required:
        displayed_score = None
        signal = "Insufficient Data"
    else:
        displayed_score = score
        if score >= 70:
            signal = "Buy"
        elif score >= 45:
            signal = "Watch"
        else:
            signal = "Sell"

    sector_rank = _rank_within(db, company, "sector") if company and not missing_required else None
    industry_rank = _rank_within(db, company, "industry") if company and not missing_required else None

    return _json_safe({
        "symbol": symbol,
        "name": company.name if company else None,
        "isin": company.isin if company else None,
        "exchange": exchange,
        "score": displayed_score,
        "signal": signal,
        "missing_required_score_categories": missing_required,
        "score_coverage_percent": coverage,
        "score_components": components,
        "score_weights": {k: round(v, 2) for k, v in normalized_weights.items()},
        "ranking_subweights": ranking_subweights,
        "rs_period_weights": rs_period_weights,
        "sector": company.sector if company else None,
        "industry": company.industry if company else None,
        "sector_rank": sector_rank,
        "industry_rank": industry_rank,
        "method_note": "Milestone-2 composite follows the client note: Fundamental 30% + Technical 25% + RS 25% + Ownership 15% + Sector 5%. Sector is included only when its real growth-ranking component is available; it is never substituted with price RS. For NSE/BSE, a score is withheld when a positively weighted fundamental or ownership category is unavailable, rather than producing a misleading partial ranking."
    })


@router.get("/ownership-details/{symbol}")
def get_ownership_details(symbol: str, exchange: str = "US"):
    try:
        provider = YahooProvider()
        return provider.get_ownership_details(symbol.upper(), exchange.upper())
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ownership details failed: {str(e)}")


@router.get("/sec-edgar/{symbol}")
def get_sec_edgar_data(
    symbol: str,
    exchange: str = "US",
    filings_limit: int = Query(12, ge=1, le=50),
):
    """Official SEC EDGAR filings + XBRL fundamental history for US stocks."""
    if exchange.upper() != "US":
        raise HTTPException(
            status_code=400,
            detail="SEC EDGAR integration applies to US-listed companies only",
        )
    try:
        data = SECFundamentalsProvider().get_snapshot(symbol.upper(), filings_limit)
        if not data.get("cik") and data.get("status") == "not_found":
            raise HTTPException(status_code=404, detail="Ticker was not found in the SEC EDGAR company list")
        # Temporary SEC/ticker-map outages return a structured 200 payload so
        # the dashboard can explain the provider problem without treating the
        # rest of the stock data as failed.
        return _json_safe(data)
    except HTTPException:
        raise
    except RuntimeError as exc:
        # Configuration/ticker-map failures are explicit, while transient SEC
        # sub-resource failures are handled inside get_snapshot as partial data.
        raise HTTPException(status_code=503, detail=str(exc))
    except Exception:
        raise HTTPException(status_code=503, detail="SEC EDGAR is temporarily unavailable. Please retry shortly.")


@router.post("/refresh-configured")
def refresh_configured_symbols_now():
    """Run the configured automatic-refresh list immediately for demonstration/testing."""
    return _json_safe(refresh_configured_market_data())


@router.get("/ranking-spec")
def get_client_ranking_spec():
    """Machine-readable Milestone-2 formulas transcribed from the client's notes."""
    return {
        "overall_composite": {
            "formula": "Fundamental*0.30 + Technical*0.25 + RS*0.25 + Ownership*0.15 + Sector*0.05",
            "weights_percent": CLIENT_COMPOSITE_WEIGHTS,
        },
        "sector_ranking": {
            "formula": "EPS_RS*0.30 + PAT_RS*0.25 + Sales_RS*0.20 + GrowthAcceleration_RS*0.15 + GrowthBreadth*0.05 + AccelerationBreadth*0.05",
            "weights_percent": CLIENT_SECTOR_WEIGHTS,
            "aggregation": "Use median stock growth for sector growth metrics rather than average.",
            "growth_breadth": "(Sales breadth + PAT breadth + EPS breadth) / 3",
            "acceleration_breadth": "EPS acceleration breadth*0.35 + PAT acceleration breadth*0.35 + Sales acceleration breadth*0.30",
        },
        "filter_layout": {
            "columns": ["Filter name", "Compare", "Value / target", "Weight", "Filter score", "Enable / disable"],
            "technical": CLIENT_TECHNICAL_FILTER_CONFIG,
            "fundamental": CLIENT_FUNDAMENTAL_FILTER_CONFIG,
            "fundamental_groups": ["EPS", "PAT", "Sales", "NPM", "CFO", "Other"],
        },
        "data_integrity": {
            "missing_values": "N/A; never fabricate a value to complete a score.",
            "ambiguous_handwritten_weights": "Remain editable/unconfirmed until the client confirms them.",
        },
    }


@router.get("/excel-feed/{symbol}")
def get_excel_feed(symbol: str, exchange: str = "US", limit: int = Query(260, ge=20, le=2000), db: Session = Depends(get_db)):
    """Refreshable JSON feed for Excel Power Query; no CSV upload is required."""
    symbol = symbol.upper()
    exchange = exchange.upper()
    company = db.query(Company).filter(Company.symbol == symbol, Company.exchange == exchange).first()
    rows = (
        db.query(OHLCV)
        .filter(OHLCV.symbol == symbol, OHLCV.exchange == exchange)
        .order_by(OHLCV.date.desc())
        .limit(limit)
        .all()
    )
    rows = list(reversed(_valid_trading_rows(rows, exchange)))
    fundamental = db.query(Fundamental).filter(Fundamental.symbol == symbol, Fundamental.exchange == exchange).first()
    ownership = db.query(Ownership).filter(Ownership.symbol == symbol, Ownership.exchange == exchange).first()
    return _json_safe({
        "symbol": symbol,
        "exchange": exchange,
        "generated_at": datetime.utcnow().isoformat() + "Z",
        "company": {
            "name": company.name if company else None,
            "sector": company.sector if company else None,
            "industry": company.industry if company else None,
        },
        "fundamental_snapshot": {
            "market_cap": fundamental.market_cap if fundamental else None,
            "trailing_eps": fundamental.trailing_eps if fundamental else None,
            "forward_eps": fundamental.forward_eps if fundamental else None,
            "revenue": fundamental.revenue if fundamental else None,
            "net_income": fundamental.net_income if fundamental else None,
            "profit_margin": fundamental.profit_margin if fundamental else None,
            "return_on_equity": fundamental.return_on_equity if fundamental else None,
            "return_on_assets": fundamental.return_on_assets if fundamental else None,
        },
        "ownership_snapshot": {
            "insider_percent": ownership.insider_percent if ownership else None,
            "institution_percent": ownership.institution_percent if ownership else None,
            "shares_outstanding": ownership.shares_outstanding if ownership else None,
            "float_shares": ownership.float_shares if ownership else None,
        },
        "ohlcv": [
            {"date": r.date.isoformat(), "open": r.open, "high": r.high, "low": r.low, "close": r.close, "volume": r.volume}
            for r in rows
        ],
        "ranking_spec": get_client_ranking_spec(),
        "excel_note": "Use Excel > Data > Get Data > From Web with this endpoint. Refresh in Excel re-requests current stored provider data.",
    })


@router.get("/excel-export/{symbol}")
def export_excel_snapshot(symbol: str, exchange: str = "US", limit: int = Query(500, ge=20, le=5000), db: Session = Depends(get_db)):
    """Download an editable workbook containing real stored data + client ranking formula sheets."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Excel export dependency is unavailable: {exc}")

    symbol = symbol.upper()
    exchange = exchange.upper()
    feed = get_excel_feed(symbol, exchange, limit, db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    summary_rows = [
        ("Symbol", symbol), ("Exchange", exchange),
        ("Company", feed["company"].get("name")),
        ("Sector", feed["company"].get("sector")),
        ("Industry", feed["company"].get("industry")),
        ("Generated At", feed.get("generated_at")),
    ]
    for row in summary_rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True)

    raw = wb.create_sheet("OHLCV")
    raw.append(["Date", "Open", "High", "Low", "Close", "Volume"])
    for cell in raw[1]: cell.font = Font(bold=True)
    for row in feed["ohlcv"]:
        raw.append([row["date"], row["open"], row["high"], row["low"], row["close"], row["volume"]])

    snap = wb.create_sheet("Fundamental_Ownership")
    snap.append(["Field", "Value"])
    for cell in snap[1]: cell.font = Font(bold=True)
    for key, value in feed["fundamental_snapshot"].items(): snap.append([f"fundamental.{key}", value])
    for key, value in feed["ownership_snapshot"].items(): snap.append([f"ownership.{key}", value])

    cfg = wb.create_sheet("Ranking_Config")
    cfg.append(["Component", "Score (0-100)", "Weight %", "Weighted Points"])
    for cell in cfg[1]: cell.font = Font(bold=True)
    order = ["fundamental", "technical", "relative_strength", "ownership", "sector"]
    for idx, key in enumerate(order, start=2):
        cfg.cell(idx, 1, key)
        cfg.cell(idx, 2, None)
        cfg.cell(idx, 3, CLIENT_COMPOSITE_WEIGHTS[key])
        cfg.cell(idx, 4, f"=IF(ISNUMBER(B{idx}),B{idx}*C{idx}/100,0)")
    cfg["A8"] = "Composite Score"
    cfg["B8"] = "=IF(COUNT(B2:B6)=0,\"\",SUM(D2:D6)/SUMPRODUCT(--ISNUMBER(B2:B6),C2:C6)*100)"
    cfg["A10"] = "Client formula"
    cfg["B10"] = "Fundamental*0.30 + Technical*0.25 + RS*0.25 + Ownership*0.15 + Sector*0.05"

    tech_filters = wb.create_sheet("Technical_Filter_Config")
    tech_filters.append(["Filter Name", "Compare", "Value / Target", "Weight", "Enabled", "Filter Score (0-100)", "Weighted Points"])
    for cell in tech_filters[1]: cell.font = Font(bold=True)
    for idx, item in enumerate(CLIENT_TECHNICAL_FILTER_CONFIG, start=2):
        tech_filters.append([item["name"], item["compare"], item["value"], item["weight"], "Yes" if item["enabled"] else "No", None, f'=IF(AND(E{idx}="Yes",ISNUMBER(F{idx})),F{idx}*D{idx},0)'])
    tech_total_row = len(CLIENT_TECHNICAL_FILTER_CONFIG) + 3
    tech_filters.cell(tech_total_row, 1, "Technical Filter Score")
    tech_filters.cell(tech_total_row, 2, f'=IFERROR(SUM(G2:G{tech_total_row-2})/SUMPRODUCT((E2:E{tech_total_row-2}="Yes")*D2:D{tech_total_row-2}),"")')

    fundamental_filters = wb.create_sheet("Fundamental_Filter_Config")
    fundamental_filters.append(["Group", "Filter Name", "Compare", "Value / Target", "Weight", "Enabled", "Filter Score (0-100)", "Weighted Points"])
    for cell in fundamental_filters[1]: cell.font = Font(bold=True)
    for idx, item in enumerate(CLIENT_FUNDAMENTAL_FILTER_CONFIG, start=2):
        fundamental_filters.append([item["group"], item["name"], item["compare"], item["value"], item["weight"], "Yes" if item["enabled"] else "No", None, f'=IF(AND(F{idx}="Yes",ISNUMBER(G{idx})),G{idx}*E{idx},0)'])
    row_cursor = len(CLIENT_FUNDAMENTAL_FILTER_CONFIG) + 3
    fundamental_filters.cell(row_cursor, 1, "Overall Fundamental Filter Score")
    fundamental_filters.cell(row_cursor, 2, f'=IFERROR(SUM(H2:H{row_cursor-2})/SUMPRODUCT((F2:F{row_cursor-2}="Yes")*E2:E{row_cursor-2}),"")')
    row_cursor += 2
    fundamental_filters.cell(row_cursor, 1, "Group Scores")
    fundamental_filters.cell(row_cursor, 1).font = Font(bold=True)
    row_cursor += 1
    for group_name in ["EPS", "PAT", "Sales", "NPM", "CFO", "Other"]:
        fundamental_filters.cell(row_cursor, 1, group_name)
        # SUMIFS keeps the workbook editable even when the client changes weights/scores.
        fundamental_filters.cell(row_cursor, 2, f'=IFERROR(SUMIFS(H$2:H${len(CLIENT_FUNDAMENTAL_FILTER_CONFIG)+1},A$2:A${len(CLIENT_FUNDAMENTAL_FILTER_CONFIG)+1},A{row_cursor})/SUMIFS(E$2:E${len(CLIENT_FUNDAMENTAL_FILTER_CONFIG)+1},A$2:A${len(CLIENT_FUNDAMENTAL_FILTER_CONFIG)+1},A{row_cursor},F$2:F${len(CLIENT_FUNDAMENTAL_FILTER_CONFIG)+1},"Yes"),"")')
        row_cursor += 1

    sector = wb.create_sheet("Sector_Ranking_Config")
    sector.append(["Sector Component", "Score (0-100)", "Weight %", "Weighted Points"])
    for cell in sector[1]: cell.font = Font(bold=True)
    for idx, (key, weight) in enumerate(CLIENT_SECTOR_WEIGHTS.items(), start=2):
        sector.cell(idx, 1, key)
        sector.cell(idx, 2, None)
        sector.cell(idx, 3, weight)
        sector.cell(idx, 4, f"=IF(ISNUMBER(B{idx}),B{idx}*C{idx}/100,0)")
    sector["A9"] = "Sector Score"
    sector["B9"] = "=IF(COUNT(B2:B7)=0,\"\",SUM(D2:D7)/SUMPRODUCT(--ISNUMBER(B2:B7),C2:C7)*100)"
    sector["A11"] = "Growth Breadth"
    sector["B11"] = "=(Sales breadth + PAT breadth + EPS breadth) / 3"
    sector["A12"] = "Acceleration Breadth"
    sector["B12"] = "=35% EPS + 35% PAT + 30% Sales acceleration breadth"

    notes = wb.create_sheet("Notes")
    notes.append(["Milestone 2 handwritten ranking notes"])
    notes["A1"].font = Font(bold=True)
    notes.append(["All market/fundamental values must come from configured providers or stored DB data; missing values stay N/A."])
    notes.append(["Ambiguous handwritten thresholds/point allocations remain editable and must not be guessed."])
    notes.append(["For a live Excel connection, use the /market/excel-feed/{symbol}?exchange=... endpoint through Power Query."])

    for sheet in wb.worksheets:
        for col in sheet.columns:
            width = min(60, max(12, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2))
            sheet.column_dimensions[col[0].column_letter].width = width

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"{symbol}_{exchange}_screener.xlsx"
    payload = stream.getvalue()
    return Response(
        content=payload,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Content-Length": str(len(payload)),
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@router.get("/provider-status")
def get_provider_status():
    """Configuration-only provider status; no third-party network call is made."""
    sec_agent = os.getenv("SEC_USER_AGENT", "").strip()
    twelve_key = os.getenv("TWELVE_DATA_API_KEY", "").strip()
    auto_symbols = os.getenv("AUTO_REFRESH_SYMBOLS", "").strip()
    return {
        "providers": {
            "yahoo_finance": {
                "configured": True,
                "uses_personal_api_key": False,
                "purpose": "US/NSE OHLCV, quotes, provider fallback fundamentals/ownership",
            },
            "sec_edgar": {
                "configured": bool(sec_agent),
                "uses_personal_api_key": False,
                "purpose": "Official US XBRL fundamentals and filing metadata",
                "required_env": "SEC_USER_AGENT",
            },
            "twelve_data_bse": {
                "configured": bool(twelve_key),
                "uses_personal_api_key": True,
                "purpose": "BSE/XBOM daily OHLCV",
                "required_env": "TWELVE_DATA_API_KEY",
            },
            "kotak_neo_india": {
                "configured": bool(os.getenv("KOTAK_NEO_API_KEY", "").strip() or os.getenv("KOTAK_NEO_ACCESS_TOKEN", "").strip()),
                "ready_for_live_quotes": bool(os.getenv("KOTAK_NEO_BASE_URL", "").strip() and os.getenv("KOTAK_NEO_ACCESS_TOKEN", "").strip()),
                "uses_personal_api_key": True,
                "purpose": "Indian-stock market data only; kept separate from US sources",
                "required_env": "KOTAK_NEO_API_KEY / KOTAK_NEO_ACCESS_TOKEN / KOTAK_NEO_BASE_URL",
                "note": "Credential/auth flow will be validated with the client's Kotak Neo account before enabling it as the live Indian provider.",
            },
        },
        "automatic_refresh": {
            "configured_symbols": [item.strip() for item in auto_symbols.split(",") if item.strip()],
            "enabled": bool(auto_symbols),
            "note": "Configure AUTO_REFRESH_SYMBOLS as EXCHANGE:SYMBOL entries; no CSV upload is required.",
        },
    }


@router.get("/india-shareholding/{symbol}")
def get_india_shareholding(symbol: str, exchange: str = "NSE", limit: int = 8):
    exchange = exchange.upper()
    if exchange not in {"NSE", "BSE"}:
        raise HTTPException(status_code=400, detail="Indian shareholding supports NSE/BSE symbols only")
    try:
        provider = IndiaShareholdingProvider()
        return _json_safe(provider.get_history(symbol.upper(), limit=limit))
    except Exception as e:
        # Keep a clean 503 instead of turning a third-party provider outage into
        # a generic 500 that makes the whole dashboard look broken.
        raise HTTPException(status_code=503, detail=f"Indian shareholding data unavailable: {str(e)}")


@router.get("/technical-summary/{symbol}")
def get_technical_summary(
    symbol: str,
    exchange: str = "US",
    timeframe: str = Query("daily", pattern="^(daily|weekly|monthly)$"),
    rs_1w_weight: float = 30,
    rs_2w_weight: float = 0,
    rs_1m_weight: float = 25,
    rs_2m_weight: float = 0,
    rs_3m_weight: float = 20,
    rs_6m_weight: float = 15,
    rs_1y_weight: float = 10,
    rs_sector_weight: float = 0,
    db: Session = Depends(get_db)
):
    symbol = symbol.upper()
    exchange = exchange.upper()

    daily_rows = (
        db.query(OHLCV)
        .filter(OHLCV.symbol == symbol, OHLCV.exchange == exchange)
        .order_by(OHLCV.date.asc())
        .all()
    )
    daily_rows = _valid_trading_rows(daily_rows, exchange)
    if len(daily_rows) < 20:
        raise HTTPException(status_code=404, detail="Not enough historical data for technical summary")

    rows = daily_rows
    if timeframe != "daily":
        import pandas as pd
        from types import SimpleNamespace

        frame = pd.DataFrame([
            {
                "date": row.date,
                "open": float(row.open),
                "high": float(row.high),
                "low": float(row.low),
                "close": float(row.close),
                "volume": float(row.volume or 0),
            }
            for row in daily_rows
        ])
        frame["date"] = pd.to_datetime(frame["date"])
        frame["actual_date"] = frame["date"]
        frame = frame.set_index("date")
        rule = "W" if timeframe == "weekly" else "ME"
        frame = frame.resample(rule).agg({
            "open": "first",
            "high": "max",
            "low": "min",
            "close": "last",
            "volume": "sum",
            "actual_date": "last",
        }).dropna(subset=["open", "high", "low", "close"])

        rows = [
            SimpleNamespace(
                date=item.actual_date.date(),
                open=float(item.open),
                high=float(item.high),
                low=float(item.low),
                close=float(item.close),
                volume=float(item.volume or 0),
            )
            for item in frame.itertuples()
        ]

    if len(rows) < 20:
        raise HTTPException(status_code=404, detail=f"Not enough {timeframe} historical data for technical summary")

    closes = [float(row.close) for row in rows]
    highs = [float(row.high) for row in rows]
    lows = [float(row.low) for row in rows]
    opens = [float(row.open) for row in rows]
    volumes = [float(row.volume or 0) for row in rows]

    ema_periods = [20, 30, 50, 100, 150, 200]
    emas = {str(period): _ema(closes, period) for period in ema_periods}
    available_emas = [emas[str(p)] for p in ema_periods if emas[str(p)] is not None]
    ema_alignment = "Unavailable"
    if len(available_emas) == len(ema_periods):
        bullish = all(available_emas[i] > available_emas[i + 1] for i in range(len(available_emas) - 1))
        bearish = all(available_emas[i] < available_emas[i + 1] for i in range(len(available_emas) - 1))
        ema_alignment = "Bullish" if bullish else "Bearish" if bearish else "Mixed"

    avg_volume_10 = sum(volumes[-10:]) / 10 if len(volumes) >= 10 else None
    avg_volume_20 = sum(volumes[-20:]) / 20
    avg_volume_40 = sum(volumes[-40:]) / 40 if len(volumes) >= 40 else None
    avg_volume_50 = sum(volumes[-50:]) / 50 if len(volumes) >= 50 else None
    volume_ratio = (volumes[-1] / avg_volume_20) if avg_volume_20 else None
    volume_ratio_50 = (volumes[-1] / avg_volume_50) if avg_volume_50 else None

    # ADR is a DAILY metric even when the chart is weekly/monthly.
    # Client/TradingView reference displays ADR(20) as an absolute price range:
    # Average(High - Low, 20 daily sessions).  Keep ADR% separately for
    # normalized comparisons and volatility logic.
    adr_daily_rows = daily_rows[-20:]
    adr_abs_values = [
        float(r.high) - float(r.low)
        for r in adr_daily_rows
        if r.high is not None and r.low is not None
    ]
    adr_percent_values = [
        ((float(r.high) - float(r.low)) / float(r.low)) * 100
        for r in adr_daily_rows
        if r.high is not None and r.low not in (None, 0)
    ]
    adr_value = sum(adr_abs_values) / len(adr_abs_values) if adr_abs_values else None
    adr_percent = sum(adr_percent_values) / len(adr_percent_values) if adr_percent_values else None

    # ATR uses Wilder's RMA (the same smoothing convention used by TradingView
    # ATR), calculated on the SELECTED timeframe.
    true_ranges = []
    for i in range(len(rows)):
        prev_close = closes[i - 1] if i > 0 else closes[i]
        true_ranges.append(max(highs[i] - lows[i], abs(highs[i] - prev_close), abs(lows[i] - prev_close)))
    atr14 = _rma(true_ranges, 14)
    atr_percent = (atr14 / closes[-1] * 100) if atr14 is not None and closes[-1] else None

    # Handwritten client ranking factors compare short ATR% averages with a
    # 20-period ATR% average.  Expose the raw comparable values so the UI can
    # score exactly those factors without inventing hidden inputs.
    atr_percent_series = [
        (true_ranges[i] / closes[i] * 100) if closes[i] else None
        for i in range(len(true_ranges))
    ]
    def _tail_average(values, period):
        usable = [v for v in values[-period:] if v is not None]
        return (sum(usable) / len(usable)) if len(usable) == period else None

    avg_atr_percent_5 = _tail_average(atr_percent_series, 5)
    avg_atr_percent_10 = _tail_average(atr_percent_series, 10)
    avg_atr_percent_20 = _tail_average(atr_percent_series, 20)
    rsi14 = _rsi(closes, 14)

    # Additional Milestone-2 technical filters from the client's handwritten notes.
    latest_price = closes[-1]
    ema20_previous = _ema(closes[:-1], 20) if len(closes) > 20 else None
    ema50_previous = _ema(closes[:-1], 50) if len(closes) > 50 else None
    ema200_previous = _ema(closes[:-1], 200) if len(closes) > 200 else None
    adx14, plus_di14, minus_di14, di_spread14 = _adx_di(highs, lows, closes, 14)
    roc_1m = _roc(closes, 21)
    roc_3m = _roc(closes, 63)
    roc_6m = _roc(closes, 126)
    roc_12m = _roc(closes, 252)
    roc_1m_two_weeks_ago = None
    if len(closes) > 31:
        historic = closes[:-10]
        roc_1m_two_weeks_ago = _roc(historic, 21)
    roc_acceleration_1m = (
        roc_1m - roc_1m_two_weeks_ago
        if roc_1m is not None and roc_1m_two_weeks_ago is not None else None
    )

    range90 = None
    if len(rows) >= 90:
        low90 = min(lows[-90:])
        high90 = max(highs[-90:])
        range90 = ((high90 - low90) / low90 * 100) if low90 else None

    recent20 = closes[-20:]
    sma20 = sum(recent20) / 20
    variance20 = sum((x - sma20) ** 2 for x in recent20) / 20
    sd20 = variance20 ** 0.5
    bb_upper = sma20 + 2 * sd20
    bb_lower = sma20 - 2 * sd20
    # Client-specified Bollinger Band width formula:
    # (Upper BB - Lower BB) * 100 / Lower BB
    bb_width = ((bb_upper - bb_lower) / bb_lower * 100) if bb_lower not in (None, 0) else None
    bb_width_20_periods_ago = None
    if len(closes) >= 40:
        old_window = closes[-40:-20]
        old_mean = sum(old_window) / 20
        old_var = sum((x - old_mean) ** 2 for x in old_window) / 20
        old_sd = old_var ** 0.5
        old_upper = old_mean + 2 * old_sd
        old_lower = old_mean - 2 * old_sd
        bb_width_20_periods_ago = ((old_upper - old_lower) / old_lower * 100) if old_lower else None

    range20 = ((max(highs[-20:]) - min(lows[-20:])) / min(lows[-20:]) * 100) if min(lows[-20:]) else None

    # Volatility trend series follows the SELECTED chart timeframe.
    # The top ADR(20D) value above intentionally remains a daily metric per the
    # earlier client rule, while this trend chart converts its candles to weekly
    # or monthly when the user changes the main chart timeframe.
    metric_rows = rows
    metric_series = []
    metric_trs = []
    running_metric_atr = None
    for i, r in enumerate(metric_rows):
        h = float(r.high)
        l = float(r.low)
        c = float(r.close)
        prev_c = float(metric_rows[i - 1].close) if i > 0 else c
        tr = max(h - l, abs(h - prev_c), abs(l - prev_c))
        metric_trs.append(tr)
        if i == 13:
            running_metric_atr = sum(metric_trs[:14]) / 14
        elif i > 13 and running_metric_atr is not None:
            running_metric_atr = ((running_metric_atr * 13) + tr) / 14

        if i < 19:
            continue

        window = metric_rows[i - 19:i + 1]
        closes20 = [float(x.close) for x in window]
        highs20 = [float(x.high) for x in window]
        lows20 = [float(x.low) for x in window]
        adr_pct_values = [
            ((float(x.high) - float(x.low)) / float(x.low)) * 100
            for x in window if float(x.low) != 0
        ]
        adr_pct_20 = (sum(adr_pct_values) / len(adr_pct_values)) if adr_pct_values else None

        mean20 = sum(closes20) / 20
        variance = sum((v - mean20) ** 2 for v in closes20) / 20
        sd = variance ** 0.5
        upper = mean20 + 2 * sd
        lower = mean20 - 2 * sd
        bb_pct = ((upper - lower) / lower * 100) if lower else None
        price_range_pct = ((max(highs20) - min(lows20)) / min(lows20) * 100) if min(lows20) else None
        atr_pct_period = (running_metric_atr / c * 100) if running_metric_atr is not None and c else None

        metric_series.append({
            "date": r.date.isoformat() if hasattr(r.date, "isoformat") else str(r.date),
            "adr_percent": round(adr_pct_20, 2) if adr_pct_20 is not None else None,
            "atr_percent": round(atr_pct_period, 2) if atr_pct_period is not None else None,
            "bollinger_width_percent": round(bb_pct, 2) if bb_pct is not None else None,
            # Keep the existing API key for frontend/backward compatibility;
            # its window is 20 periods of the selected timeframe.
            "range_20d_percent": round(price_range_pct, 2) if price_range_pct is not None else None,
        })
    periods_per_52w = 252 if timeframe == "daily" else 52 if timeframe == "weekly" else 12
    lookback_52w = min(periods_per_52w, len(rows))
    high_52w = max(highs[-lookback_52w:])
    distance_52w_high = ((high_52w - closes[-1]) / high_52w * 100) if high_52w else None
    high_52w_slice = highs[-lookback_52w:]
    high_52w_index = max(range(len(high_52w_slice)), key=lambda i: high_52w_slice[i]) if high_52w_slice else None
    periods_since_52w_high = (len(high_52w_slice) - 1 - high_52w_index) if high_52w_index is not None else None

    # VCP/consolidation detection. Three successive 20-period windows are used
    # as transparent contractions. When both price depth and ATR% contract in
    # sequence, the final contraction high becomes the pivot (client-defined).
    contractions = []
    final_contraction_high = None
    if len(rows) >= 60:
        for start_i, end_i in [(-60, -40), (-40, -20), (-20, None)]:
            segment = rows[start_i:end_i]
            seg_high = max(float(r.high) for r in segment)
            seg_low = min(float(r.low) for r in segment)
            depth = ((seg_high - seg_low) / seg_high * 100) if seg_high else None
            seg_tr = []
            for j, r in enumerate(segment):
                prev = float(segment[j-1].close) if j > 0 else float(r.close)
                seg_tr.append(max(float(r.high)-float(r.low), abs(float(r.high)-prev), abs(float(r.low)-prev)))
            seg_atr = _rma(seg_tr, 14) if len(seg_tr) >= 14 else (sum(seg_tr) / len(seg_tr) if seg_tr else None)
            seg_atr_pct = (seg_atr / float(segment[-1].close) * 100) if seg_atr and segment[-1].close else None
            seg_closes = [float(r.close) for r in segment if r.close is not None]
            seg_mean = (sum(seg_closes) / len(seg_closes)) if seg_closes else None
            seg_std = (sum((x - seg_mean) ** 2 for x in seg_closes) / len(seg_closes)) ** 0.5 if seg_closes and seg_mean else None
            seg_std_pct = (seg_std / seg_mean * 100) if seg_std is not None and seg_mean else None
            seg_avg_volume = (sum(float(r.volume or 0) for r in segment) / len(segment)) if segment else None
            contractions.append({
                "depth_percent": depth,
                "atr_percent": seg_atr_pct,
                "standard_deviation_percent": seg_std_pct,
                "average_volume": seg_avg_volume,
                "high": seg_high,
                "low": seg_low,
            })

    vcp_stage = "Not Detected"
    if len(contractions) == 3:
        depths = [c["depth_percent"] for c in contractions]
        atrs = [c["atr_percent"] for c in contractions]
        stds = [c["standard_deviation_percent"] for c in contractions]
        vols = [c["average_volume"] for c in contractions]
        price_contracting = all(v is not None for v in depths) and depths[1] < depths[0] and depths[2] < depths[1]
        atr_contracting = all(v is not None for v in atrs) and atrs[1] < atrs[0] and atrs[2] < atrs[1]
        std_contracting = all(v is not None for v in stds) and stds[1] < stds[0] and stds[2] < stds[1]
        volume_contracting = all(v is not None for v in vols) and vols[1] < vols[0] and vols[2] < vols[1]
        if price_contracting and atr_contracting and std_contracting and volume_contracting:
            vcp_stage = "VCP Contraction"
            final_contraction_high = contractions[-1]["high"]

    # If a formal VCP is not detected, use the recent consolidation high as the
    # fallback pivot. Never use today's high itself as the pivot.
    consolidation_rows = rows[-21:-1] if len(rows) >= 21 else rows[:-1]
    consolidation_high = max((float(r.high) for r in consolidation_rows), default=None)
    pivot = final_contraction_high if final_contraction_high is not None else consolidation_high

    buffer = 0.003
    latest_close = closes[-1]
    latest_open = opens[-1]
    latest_midpoint = (highs[-1] + lows[-1]) / 2

    breakout_status = "Unavailable"
    breakout_strength = None
    near_pivot = False
    if pivot is not None:
        near_pivot = (pivot * 0.95) <= latest_close <= (pivot * 1.02)
        above_pivot = latest_close > pivot * (1 + buffer)
        price_confirmation = latest_close > latest_open and latest_close >= latest_midpoint
        volume_confirmation = volume_ratio_50 is not None and volume_ratio_50 >= 1.4

        if above_pivot and price_confirmation and volume_confirmation:
            breakout_status = "Confirmed Breakout"
        elif latest_close > pivot:
            breakout_status = "Potential Breakout"
        elif near_pivot:
            breakout_status = "Near Pivot"
        else:
            breakout_status = "No Breakout"

        # Client-requested 0-100 breakout-strength model.
        points = 0
        if latest_close > pivot: points += 20
        if volume_ratio_50 is not None and volume_ratio_50 >= 1.4: points += 20
        if volume_ratio_50 is not None and volume_ratio_50 >= 2.0: points += 10
        if latest_close > pivot * 1.005: points += 10
        if latest_close > pivot * 1.01: points += 10

        atr_contraction = False
        if len(contractions) >= 2:
            a = contractions[-2].get("atr_percent")
            b = contractions[-1].get("atr_percent")
            atr_contraction = a is not None and b is not None and b < a
        if atr_contraction: points += 10

        if len(rows) >= 10:
            tight_high = max(highs[-10:])
            tight_low = min(lows[-10:])
            tight_range = ((tight_high - tight_low) / tight_high * 100) if tight_high else None
            if tight_range is not None and tight_range <= 10:
                points += 10

        breakout_strength = min(100, points)

    gap_percent = None
    gap_classification = "Unavailable"
    if len(rows) >= 2 and closes[-2]:
        gap_percent = ((opens[-1] - closes[-2]) / closes[-2]) * 100
        if gap_percent >= 8:
            gap_classification = "Excessive Gap Breakout"
        elif gap_percent >= 2:
            gap_classification = "Gap Breakout"
        else:
            gap_classification = "Normal / No Material Gap"

    if breakout_status == "Confirmed Breakout":
        pattern = "Confirmed Breakout"
    elif vcp_stage == "VCP Contraction":
        pattern = "VCP / Volatility Contraction"
    elif range20 is not None and range20 <= 10:
        pattern = "Tight Consolidation"
    elif near_pivot:
        pattern = "Near Pivot"
    else:
        pattern = "None"

    rs_period_weights = {"1w": rs_1w_weight, "2w": rs_2w_weight, "1m": rs_1m_weight, "2m": rs_2m_weight, "3m": rs_3m_weight, "6m": rs_6m_weight, "1y": rs_1y_weight, "sector": rs_sector_weight}
    rs_rating, rs_weighted_relative_return, rs_metrics, benchmark_name, rs_chart = _weighted_rs_against_benchmark(daily_rows, exchange, rs_period_weights, db=db)

    # Client handwritten comparison table: raw returns stay separate from RS scoring weights.
    stock_points = [
        (r.date, float(r.open) if r.open is not None else float(r.close), float(r.close))
        for r in daily_rows if r.date is not None and r.close is not None
    ]
    company = db.query(Company).filter(Company.symbol == symbol, Company.exchange == exchange).first()
    if exchange == "US":
        index_rows = [
            {"key": "dow_jones", "label": "Dow Jones", "returns": _yf_period_returns("^DJI")},
            {"key": "sp500", "label": "S&P 500", "returns": _yf_period_returns("^GSPC")},
        ]
    else:
        index_rows = [
            {"key": "nifty50", "label": "NIFTY 50", "returns": _yf_period_returns("^NSEI")},
            {"key": "nifty500", "label": "NIFTY 500", "returns": _yf_period_returns("^CRSLDX")},
        ]
    industry_returns = _peer_group_period_returns(
        db, exchange, "industry", company.industry if company else None,
        exclude_symbol=symbol, minimum_peers=5
    )
    sector_returns = _peer_group_period_returns(
        db, exchange, "sector", company.sector if company else None,
        exclude_symbol=symbol, minimum_peers=5
    )
    rs_comparison = {
        "periods": ["1d", "1w", "2w", "3w", "1m", "2m", "3m", "6m", "1y"],
        "rows": [
            {"key": "stock", "label": "Stock Return", "returns": _period_returns_from_points(stock_points)},
            *index_rows,
            {
                "key": "industry", "label": "Industry",
                "name": company.industry if company else None,
                "returns": industry_returns,
                "note": None if industry_returns else "Insufficient stored peer history; N/A shown.",
            },
            {
                "key": "sector", "label": "Sector",
                "name": company.sector if company else None,
                "returns": sector_returns,
                "note": None if sector_returns else "Insufficient stored peer history; N/A shown.",
            },
        ],
        "method": "Current-period open to latest close; industry/sector are equal-weight averages of at least 5 stored peers, excluding the selected stock. If fewer peers are available, N/A is shown. These raw returns do not use RS scoring weights.",
    }

    delivery_summary = _nse_delivery_summary(symbol) if exchange == "NSE" else {
        "available": False,
        "source": exchange,
        "note": "True delivery percentage requires exchange deliverable-quantity data; ordinary OHLCV volume is not substituted.",
    }

    client_technical_filters = {
        "trend": {
            "price": latest_price,
            "price_gt_ema20": (latest_price > emas.get("20")) if emas.get("20") is not None else None,
            "ema20_slope_positive": (emas.get("20") > ema20_previous) if emas.get("20") is not None and ema20_previous is not None else None,
            "price_gt_ema50": (latest_price > emas.get("50")) if emas.get("50") is not None else None,
            "ema50_slope_positive": (emas.get("50") > ema50_previous) if emas.get("50") is not None and ema50_previous is not None else None,
            "price_gt_ema200": (latest_price > emas.get("200")) if emas.get("200") is not None else None,
            "ema200_slope_positive": (emas.get("200") > ema200_previous) if emas.get("200") is not None and ema200_previous is not None else None,
        },
        "strength": {
            "rs_score": rs_rating,
            "distance_from_52w_high_percent": distance_52w_high,
            "periods_since_52w_high": periods_since_52w_high,
            "new_52w_high_within_30_periods": (periods_since_52w_high <= 30) if periods_since_52w_high is not None else None,
            "adx_14": adx14,
            "plus_di_14": plus_di14,
            "minus_di_14": minus_di14,
            "di_spread": di_spread14,
            "di_spread_gt_10": (di_spread14 > 10) if di_spread14 is not None else None,
        },
        "momentum": {
            "roc_1m_percent": roc_1m,
            "roc_3m_percent": roc_3m,
            "roc_6m_percent": roc_6m,
            "roc_12m_percent": roc_12m,
            "roc_1m_two_weeks_ago_percent": roc_1m_two_weeks_ago,
            "roc_1m_acceleration_points": roc_acceleration_1m,
        },
        "participation": {
            "average_volume_10": avg_volume_10,
            "average_volume_20": avg_volume_20,
            "average_volume_40": avg_volume_40,
            "average_volume_50": avg_volume_50,
            "delivery_percent": delivery_summary,
        },
        "volatility": {
            "atr_percent": atr_percent,
            "average_atr_percent_5": avg_atr_percent_5,
            "average_atr_percent_10": avg_atr_percent_10,
            "average_atr_percent_20": avg_atr_percent_20,
            "bollinger_width_percent": bb_width,
            "bollinger_width_20_periods_ago_percent": bb_width_20_periods_ago,
            "bb_width_contracting": (bb_width < bb_width_20_periods_ago) if bb_width is not None and bb_width_20_periods_ago is not None else None,
        },
        "base_formation": {
            "range_20_period_percent": range20,
            "range_90_period_percent": range90,
            "range_90_le_20_percent": (range90 <= 20) if range90 is not None else None,
            "vcp_stage": vcp_stage,
            "standard_deviation_contracting": std_contracting if len(contractions) == 3 else None,
            "volume_contracting": volume_contracting if len(contractions) == 3 else None,
        },
        "note": "These are real-data filter inputs from the client handwritten technical sheet. Ambiguous handwritten point allocations remain configurable rather than guessed.",
    }

    return _json_safe({
        "symbol": symbol,
        "exchange": exchange,
        "timeframe": timeframe,
        "ema": {key: (round(value, 2) if value is not None else None) for key, value in emas.items()},
        "ema_alignment": ema_alignment,
        "average_volume_10": round(avg_volume_10, 2) if avg_volume_10 is not None else None,
        "average_volume_20": round(avg_volume_20, 2),
        "average_volume_40": round(avg_volume_40, 2) if avg_volume_40 is not None else None,
        "average_volume_50": round(avg_volume_50, 2) if avg_volume_50 is not None else None,
        "volume_ratio": round(volume_ratio, 2) if volume_ratio is not None else None,
        "volume_ratio_50": round(volume_ratio_50, 2) if volume_ratio_50 is not None else None,
        "adr_20": round(adr_value, 2) if adr_value is not None else None,
        "adr_percent": round(adr_percent, 2) if adr_percent is not None else None,
        "atr_14": round(atr14, 2) if atr14 is not None else None,
        "atr_percent": round(atr_percent, 2) if atr_percent is not None else None,
        "average_atr_percent_5": round(avg_atr_percent_5, 2) if avg_atr_percent_5 is not None else None,
        "average_atr_percent_10": round(avg_atr_percent_10, 2) if avg_atr_percent_10 is not None else None,
        "average_atr_percent_20": round(avg_atr_percent_20, 2) if avg_atr_percent_20 is not None else None,
        "rsi_14": round(rsi14, 2) if rsi14 is not None else None,
        "bollinger_width_percent": round(bb_width, 2) if bb_width is not None else None,
        "range_20d_percent": round(range20, 2) if range20 is not None else None,
        "technical_metric_series": metric_series[-260:],
        "technical_metric_timeframe": timeframe,
        "technical_metric_window_periods": 20,
        "distance_from_52w_high_percent": round(distance_52w_high, 2) if distance_52w_high is not None else None,
        "pivot": round(pivot, 2) if pivot is not None else None,
        "breakout_status": breakout_status,
        "breakout_strength": breakout_strength,
        "gap_percent": round(gap_percent, 2) if gap_percent is not None else None,
        "gap_classification": gap_classification,
        "vcp_stage": vcp_stage,
        "vcp_contractions": [{k: (round(v, 2) if isinstance(v, (int, float)) and v is not None else v) for k, v in item.items()} for item in contractions],
        "vcp_standard_deviation_contraction": (std_contracting if len(contractions) == 3 else None),
        "vcp_volume_contraction": (volume_contracting if len(contractions) == 3 else None),
        "pattern": pattern,
        "client_technical_filters": client_technical_filters,
        "rs_comparison": rs_comparison,
        "volume_delivery": delivery_summary,
        "rs_rating": rs_rating,
        "rs_available": rs_rating is not None and len(rs_chart) > 1,
        "rs_universe": {
            "market": _rs_market_group(exchange),
            "market_exchanges": list(_rs_market_exchanges(exchange)),
            "target_size": _rs_percentile_total(exchange),
            "scored_stocks_available": int(rs_metrics.get("_scored_stocks_available", 0) or 0),
            "coverage_percent": round(
                min(100.0, (float(rs_metrics.get("_scored_stocks_available", 0) or 0) / _rs_percentile_total(exchange)) * 100.0),
                2,
            ),
            "complete": int(rs_metrics.get("_scored_stocks_available", 0) or 0) >= _rs_percentile_total(exchange),
            "note": (
                f"Client percentile denominator is fixed at {_rs_percentile_total(exchange):,} for {_rs_market_group(exchange)}. "
                f"The RS score is provisional until {_rs_percentile_total(exchange):,} stored stocks have usable comparison history."
                if int(rs_metrics.get("_scored_stocks_available", 0) or 0) < _rs_percentile_total(exchange)
                else f"The stored RS comparison universe meets the client-required {_rs_percentile_total(exchange):,}-stock denominator for {_rs_market_group(exchange)}."
            ),
        },
        "rs_weighted_relative_return_percent": round(rs_weighted_relative_return, 2) if rs_weighted_relative_return is not None else None,
        "rs_benchmark": benchmark_name,
        "rs_periods": rs_metrics,
        "rs_chart": rs_chart,
        "rs_period_weights": rs_period_weights,
        "rs_formula": {
            "period_return": "((current close / current period candle open) - 1) * 100",
            "relative_return": "stock return % - benchmark return %",
            "relative_return_uses_editable_weights": False,
            "stock_percentile": f"((lower stocks + 0.5 * equal stocks) * 100) / {_rs_percentile_total(exchange)}",
            "stock_percentile_denominator": _rs_percentile_total(exchange),
            "final_rs_score": "sum(period percentile * enabled weight) / sum(enabled weights)",
            "default_weights": {"1w": 30, "1m": 25, "3m": 20, "6m": 15, "1y": 10, "2w": 0, "2m": 0, "sector": 0},
        },
        "rs_note": f"Relative Strength: period relative returns are raw market calculations and never change when RS score weights change. Each period uses TradingView-style current period candle return (period open to current close); relative return = stock return % - benchmark return %. Stock percentiles use the client-required fixed denominator of {_rs_percentile_total(exchange):,} stocks for {_rs_market_group(exchange)}: [(lower stocks + 0.5 x equal stocks) x 100 / {_rs_percentile_total(exchange)}]. Final RS Score uses weighted percentile components. Default weights are 1W x 30% + 1M x 25% + 3M x 20% + 6M x 15% + 12M x 10%. 2W/2M and Sector RS are optional components with zero default weight; Sector RS is included only when its weight is greater than 0.",
        "criteria_note": f"Metrics use the selected {timeframe} timeframe. For a detected VCP, Pivot = highest high of the final contraction; otherwise it is the recent consolidation high. Near Pivot = 95%-102% of pivot. Confirmed breakout requires close > pivot by 0.3%, volume >= 1.4x 50-period average, close > open, and close in the upper half of the period's range. VCP requires successive price-depth, ATR%, standard-deviation, and average-volume contractions."
    })


@router.get("/benchmark/{exchange}")
def get_benchmark(exchange: str, limit: int = 400):
    exchange = exchange.upper()
    if exchange == "US":
        ticker_symbol = "^GSPC"
        benchmark_name = "S&P 500"
    elif exchange in ["NSE", "BSE"]:
        ticker_symbol = "^CRSLDX"
        benchmark_name = "NIFTY 500"
    else:
        raise HTTPException(status_code=400, detail="Unsupported exchange")

    # Yahoo is used first because it exposes both requested index series.
    try:
        data = yf.Ticker(ticker_symbol).history(period="5y", interval="1d", auto_adjust=False)
        if data is not None and not data.empty:
            values = [
                {"date": idx.date().isoformat(), "close": float(row["Close"])}
                for idx, row in data.iterrows()
                if row.get("Close") is not None
            ]
            return {
                "name": benchmark_name,
                "symbol": ticker_symbol,
                "data": values[-limit:],
                "method": "Broad-market index history from Yahoo Finance",
            }
    except Exception:
        pass

    # US fallback for deployments where Yahoo index history is temporarily unavailable.
    if exchange == "US":
        api_key = os.getenv("TWELVE_DATA_API_KEY")
        if api_key:
            try:
                response = requests.get(
                    "https://api.twelvedata.com/time_series",
                    params={"symbol": "SPY", "interval": "1day", "outputsize": limit, "apikey": api_key},
                    timeout=20,
                )
                payload = response.json()
                if payload.get("status") != "error":
                    values = payload.get("values", [])
                    return {
                        "name": "S&P 500 (SPY fallback)",
                        "symbol": "SPY",
                        "data": [{"date": item["datetime"], "close": float(item["close"])} for item in reversed(values)],
                        "warning": "Using SPY ETF as fallback because the S&P 500 index feed was unavailable.",
                    }
            except Exception:
                pass

    return {
        "name": benchmark_name,
        "symbol": ticker_symbol,
        "data": [],
        "warning": f"{benchmark_name} data is temporarily unavailable from the configured providers.",
    }


@router.get("/chart/{symbol}")
def get_chart_data(
    symbol: str,
    exchange: str = "US",
    timeframe: str = Query("daily", pattern="^(daily|weekly|monthly)$"),
    limit: int = 260,
    db: Session = Depends(get_db)
):
    rows = (
        db.query(OHLCV)
        .filter(
            OHLCV.symbol == symbol.upper(),
            OHLCV.exchange == exchange.upper()
        )
        .order_by(OHLCV.date.asc())
        .all()
    )
    rows = _valid_trading_rows(rows, exchange)

    if not rows:
        raise HTTPException(
            status_code=404,
            detail="No stored data found"
        )

    data = [
        {
            "date": row.date,
            "open": row.open,
            "high": row.high,
            "low": row.low,
            "close": row.close,
            "volume": row.volume
        }
        for row in rows
    ]

    if timeframe == "daily":
        result = data

    else:
        import pandas as pd

        df = pd.DataFrame(data)
        df["date"] = pd.to_datetime(df["date"])
        df["actual_date"] = df["date"]
        df = df.set_index("date")

        rule = "W" if timeframe == "weekly" else "ME"

        df = df.resample(rule).agg({
            "open": "first",
            "high": "max",
            "low": "min",
            "close": "last",
            "volume": "sum",
            "actual_date": "last"
        }).dropna(subset=["open", "high", "low", "close"])

        df = df.reset_index(drop=True).rename(columns={"actual_date": "date"})

        result = df.to_dict(orient="records")

    return _json_safe({
        "symbol": symbol.upper(),
        "exchange": exchange.upper(),
        "timeframe": timeframe,
        "count": len(result[-limit:]),
        "data": result[-limit:]
    })

@router.post("/fundamentals/{symbol}")
def refresh_fundamentals(
    symbol: str,
    exchange: str = "US",
    db: Session = Depends(get_db)
):
    try:
        exchange = exchange.upper()
        if exchange not in {"US", "NSE", "BSE"}:
            raise HTTPException(status_code=400, detail="Unsupported exchange")

        provider = YahooProvider()
        data = provider.get_fundamentals(symbol.upper(), exchange)

        # Use statement-derived ratios for ROE/ROA when available. This keeps
        # the snapshot consistent with the annual table and uses average balance
        # sheet denominators instead of an opaque provider summary ratio.
        try:
            history = provider.get_fundamental_history(symbol.upper(), exchange)
            latest_annual = (history.get("annual") or [None])[0]
            if latest_annual:
                if latest_annual.get("roe") is not None:
                    data["return_on_equity"] = latest_annual["roe"] / 100
                if latest_annual.get("roa") is not None:
                    data["return_on_assets"] = latest_annual["roa"] / 100
        except Exception:
            pass

        result = sync_fundamental_data(
            db=db,
            symbol=symbol.upper(),
            exchange=exchange,
            data=data
        )

        return {
            **result,
            "fundamentals": data
        }

    except HTTPException:
        raise

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Fundamental refresh failed: {str(e)}"
        )

@router.get("/fundamentals/{symbol}")
def get_fundamentals(
    symbol: str,
    exchange: str = "US",
    db: Session = Depends(get_db)
):
    fundamental = (
        db.query(Fundamental)
        .filter(
            Fundamental.symbol == symbol.upper(),
            Fundamental.exchange == exchange.upper()
        )
        .first()
    )

    ownership = (
        db.query(Ownership)
        .filter(
            Ownership.symbol == symbol.upper(),
            Ownership.exchange == exchange.upper()
        )
        .first()
    )

    company = (
        db.query(Company)
        .filter(Company.symbol == symbol.upper(), Company.exchange == exchange.upper())
        .first()
    )

    if not fundamental and not ownership:
        raise HTTPException(
            status_code=404,
            detail="No fundamental data found"
        )

    return {
        "symbol": symbol.upper(),
        "name": company.name if company else None,
        "isin": company.isin if company else None,
        "exchange": exchange.upper(),
        "fundamentals": {
            "market_cap": fundamental.market_cap if fundamental else None,
            "trailing_eps": fundamental.trailing_eps if fundamental else None,
            "forward_eps": fundamental.forward_eps if fundamental else None,
            "revenue": fundamental.revenue if fundamental else None,
            "net_income": fundamental.net_income if fundamental else None,
            "profit_margin": fundamental.profit_margin if fundamental else None,
            "return_on_equity": fundamental.return_on_equity if fundamental else None,
            "return_on_assets": fundamental.return_on_assets if fundamental else None,
        },
        "ownership": {
            "insider_percent": ownership.insider_percent if ownership else None,
            "institution_percent": ownership.institution_percent if ownership else None,
            "shares_outstanding": ownership.shares_outstanding if ownership else None,
            "float_shares": ownership.float_shares if ownership else None,
        }
    }

@router.get("/fundamentals-history/{symbol}")
def get_fundamentals_history(
    symbol: str,
    exchange: str = "US"
):
    try:
        exchange = exchange.upper()
        if exchange not in {"US", "NSE", "BSE"}:
            raise HTTPException(status_code=400, detail="Unsupported exchange")

        provider = YahooProvider()

        data = provider.get_fundamental_history(
            symbol.upper(), exchange
        )

        return data

    except HTTPException:
        raise

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Fundamental history failed: {str(e)}"
        )

@router.get("/indicators/{symbol}")
def get_indicators(
    symbol: str,
    exchange: str = "US",
    timeframe: str = "daily",
    sma_short: int = 20,
    sma_long: int = 50,
    rsi_period: int = 14,
    db: Session = Depends(get_db)
):
    rows = (
        db.query(OHLCV)
        .filter(
            OHLCV.symbol == symbol.upper(),
            OHLCV.exchange == exchange.upper()
        )
        .order_by(OHLCV.date.asc())
        .all()
    )
    rows = _valid_trading_rows(rows, exchange)

    if not rows:
        raise HTTPException(
            status_code=404,
            detail="No historical data found"
        )

    timeframe = timeframe.lower()

    # Build closes for selected timeframe
    if timeframe == "daily":
        closes = [float(r.close) for r in rows]

    elif timeframe == "weekly":
        weekly = {}

        for r in rows:
            # ISO year + ISO week prevents week-number collisions across years
            iso = r.date.isocalendar()
            key = (iso.year, iso.week)
            weekly[key] = float(r.close)

        closes = list(weekly.values())

    elif timeframe == "monthly":
        monthly = {}

        for r in rows:
            key = (r.date.year, r.date.month)
            monthly[key] = float(r.close)

        closes = list(monthly.values())

    else:
        raise HTTPException(
            status_code=400,
            detail="Timeframe must be daily, weekly, or monthly"
        )

    minimum_required = max(
        sma_short,
        sma_long,
        rsi_period + 1
    )

    if len(closes) < minimum_required:
        raise HTTPException(
            status_code=404,
            detail=f"Not enough {timeframe} historical data for selected indicator settings"
        )

    # SMA
    sma_short_value = sum(closes[-sma_short:]) / sma_short
    sma_long_value = sum(closes[-sma_long:]) / sma_long

    # EMA
    def calculate_ema(values, period):
        multiplier = 2 / (period + 1)

        ema = sum(values[:period]) / period

        for price in values[period:]:
            ema = ((price - ema) * multiplier) + ema

        return ema

    ema_short_value = calculate_ema(closes, sma_short)
    ema_long_value = calculate_ema(closes, sma_long)
    ema_150_value = calculate_ema(closes, 150) if len(closes) >= 150 else None
    ema_200_value = calculate_ema(closes, 200) if len(closes) >= 200 else None

    # RSI
    changes = [
        closes[i] - closes[i - 1]
        for i in range(1, len(closes))
    ]

    gains = [max(change, 0) for change in changes]
    losses = [max(-change, 0) for change in changes]

    def calculate_rma(values, period):
        rma = sum(values[:period]) / period
        alpha = 1 / period

        for value in values[period:]:
            rma = (alpha * value) + ((1 - alpha) * rma)

        return rma

    avg_gain = calculate_rma(gains, rsi_period)
    avg_loss = calculate_rma(losses, rsi_period)

    if avg_loss == 0:
        rsi_value = 100
    else:
        rs = avg_gain / avg_loss
        rsi_value = 100 - (100 / (1 + rs))

    return {
        "symbol": symbol.upper(),
        "exchange": exchange.upper(),
        "timeframe": timeframe,
        "settings": {
            "sma_short": sma_short,
            "sma_long": sma_long,
            "rsi_period": rsi_period
        },
        "sma_short": round(sma_short_value, 2),
        "sma_long": round(sma_long_value, 2),
        "ema_short": round(ema_short_value, 2),
        "ema_long": round(ema_long_value, 2),
        "ema_150": round(ema_150_value, 2) if ema_150_value is not None else None,
        "ema_200": round(ema_200_value, 2) if ema_200_value is not None else None,
        "rsi": round(rsi_value, 2)
    }

# ---------------------------------------------------------------------------
# Universe screener: stored-data-first list view + filtered Excel export.
# This powers the client's requested "many stocks in one table" workflow.
# It intentionally uses only data already stored in this application's DB.
# Missing fields stay N/A; no synthetic fundamentals/ownership values are made.
# ---------------------------------------------------------------------------

SCREENER_COLUMN_LABELS = {
    "symbol": "Symbol",
    "name": "Company",
    "isin": "ISIN",
    "exchange": "Exchange",
    "sector": "Sector",
    "industry": "Industry",
    "close": "LTP",
    "volume": "Volume",
    "market_cap": "Market Cap",
    "trailing_eps": "EPS",
    "forward_eps": "Forward EPS",
    "revenue": "Revenue",
    "net_income": "Net Income",
    "profit_margin": "Profit Margin %",
    "return_on_equity": "ROE %",
    "return_on_assets": "ROA %",
    "institution_percent": "Institution %",
    "insider_percent": "Insider %",
    "shares_outstanding": "Shares Outstanding",
    "float_shares": "Float Shares",
    "distance_52w_high": "Distance From 52W High %",
    "distance_52w_low": "Distance From 52W Low %",
    "volume_ratio": "Volume / 52W Avg",
    "latest_date": "Latest Price Date",
    "data_coverage": "Data Coverage %",
}

SCREENER_DEFAULT_COLUMNS = [
    "symbol", "name", "isin", "exchange", "close", "market_cap", "trailing_eps",
    "profit_margin", "return_on_equity", "institution_percent",
    "distance_52w_high", "volume_ratio", "data_coverage",
]


def _screener_market_exchanges(market: str):
    market = (market or "ALL").upper()
    if market == "US":
        return ["US"]
    if market in {"INDIA", "IN"}:
        return ["NSE", "BSE"]
    if market == "NSE":
        return ["NSE"]
    if market == "BSE":
        return ["BSE"]
    return ["US", "NSE", "BSE"]


def _screener_query_parts(db: Session, market: str):
    exchanges = _screener_market_exchanges(market)
    latest_dates = (
        db.query(
            OHLCV.symbol.label("symbol"),
            OHLCV.exchange.label("exchange"),
            func.max(OHLCV.date).label("max_date"),
        )
        .filter(OHLCV.exchange.in_(exchanges))
        .group_by(OHLCV.symbol, OHLCV.exchange)
        .subquery()
    )
    latest = aliased(OHLCV)
    cutoff = date.today() - timedelta(days=370)
    year_stats = (
        db.query(
            OHLCV.symbol.label("symbol"),
            OHLCV.exchange.label("exchange"),
            func.max(OHLCV.high).label("high_52w"),
            func.min(OHLCV.low).label("low_52w"),
            func.avg(OHLCV.volume).label("avg_volume_52w"),
        )
        .filter(OHLCV.exchange.in_(exchanges), OHLCV.date >= cutoff)
        .group_by(OHLCV.symbol, OHLCV.exchange)
        .subquery()
    )

    distance_high = case(
        (year_stats.c.high_52w > 0, ((year_stats.c.high_52w - latest.close) / year_stats.c.high_52w) * 100.0),
        else_=None,
    )
    distance_low = case(
        (year_stats.c.low_52w > 0, ((latest.close - year_stats.c.low_52w) / year_stats.c.low_52w) * 100.0),
        else_=None,
    )
    volume_ratio = case(
        (year_stats.c.avg_volume_52w > 0, latest.volume / year_stats.c.avg_volume_52w),
        else_=None,
    )

    # Prefer rows with useful real data instead of filling the first page with
    # N/A-only listings.  This is a completeness indicator only; it never
    # substitutes or estimates a missing value.
    data_points = (
        case((latest.close.isnot(None), 1), else_=0)
        + case((Fundamental.market_cap.isnot(None), 1), else_=0)
        + case((Fundamental.trailing_eps.isnot(None), 1), else_=0)
        + case((Fundamental.revenue.isnot(None), 1), else_=0)
        + case((Fundamental.net_income.isnot(None), 1), else_=0)
        + case((Fundamental.profit_margin.isnot(None), 1), else_=0)
        + case((Fundamental.return_on_equity.isnot(None), 1), else_=0)
        + case((Ownership.institution_percent.isnot(None), 1), else_=0)
        + case((or_(Company.sector.isnot(None), Company.industry.isnot(None)), 1), else_=0)
    )
    data_coverage = (data_points * 100.0 / 9.0)

    query = (
        db.query(
            Company.symbol.label("symbol"),
            Company.name.label("name"),
            Company.isin.label("isin"),
            Company.exchange.label("exchange"),
            Company.sector.label("sector"),
            Company.industry.label("industry"),
            latest.close.label("close"),
            latest.volume.label("volume"),
            latest.date.label("latest_date"),
            Fundamental.market_cap.label("market_cap"),
            Fundamental.trailing_eps.label("trailing_eps"),
            Fundamental.forward_eps.label("forward_eps"),
            Fundamental.revenue.label("revenue"),
            Fundamental.net_income.label("net_income"),
            Fundamental.profit_margin.label("profit_margin"),
            Fundamental.return_on_equity.label("return_on_equity"),
            Fundamental.return_on_assets.label("return_on_assets"),
            Ownership.institution_percent.label("institution_percent"),
            Ownership.insider_percent.label("insider_percent"),
            Ownership.shares_outstanding.label("shares_outstanding"),
            Ownership.float_shares.label("float_shares"),
            distance_high.label("distance_52w_high"),
            distance_low.label("distance_52w_low"),
            volume_ratio.label("volume_ratio"),
            data_coverage.label("data_coverage"),
        )
        .outerjoin(Fundamental, (Fundamental.symbol == Company.symbol) & (Fundamental.exchange == Company.exchange))
        .outerjoin(Ownership, (Ownership.symbol == Company.symbol) & (Ownership.exchange == Company.exchange))
        .outerjoin(latest_dates, (latest_dates.c.symbol == Company.symbol) & (latest_dates.c.exchange == Company.exchange))
        .outerjoin(latest, (latest.symbol == Company.symbol) & (latest.exchange == Company.exchange) & (latest.date == latest_dates.c.max_date))
        .outerjoin(year_stats, (year_stats.c.symbol == Company.symbol) & (year_stats.c.exchange == Company.exchange))
        .filter(Company.is_active == 1, Company.exchange.in_(exchanges))
    )
    # Defensive query-time equity filter: even if a provider sync is delayed,
    # legacy US warrants/units/SPAC rows cannot leak into the normal screener.
    query = apply_eligible_equity_filter(query, exchanges)

    expressions = {
        "symbol": Company.symbol,
        "name": Company.name,
        "isin": Company.isin,
        "exchange": Company.exchange,
        "sector": Company.sector,
        "industry": Company.industry,
        "close": latest.close,
        "volume": latest.volume,
        "market_cap": Fundamental.market_cap,
        "trailing_eps": Fundamental.trailing_eps,
        "forward_eps": Fundamental.forward_eps,
        "revenue": Fundamental.revenue,
        "net_income": Fundamental.net_income,
        "profit_margin": Fundamental.profit_margin,
        "return_on_equity": Fundamental.return_on_equity,
        "return_on_assets": Fundamental.return_on_assets,
        "institution_percent": Ownership.institution_percent,
        "insider_percent": Ownership.insider_percent,
        "shares_outstanding": Ownership.shares_outstanding,
        "float_shares": Ownership.float_shares,
        "distance_52w_high": distance_high,
        "distance_52w_low": distance_low,
        "volume_ratio": volume_ratio,
        "latest_date": latest.date,
        "data_coverage": data_coverage,
    }
    return query, expressions


def _apply_screener_filters(
    query,
    expressions,
    q=None,
    sector=None,
    industry=None,
    market_cap_min=None,
    market_cap_max=None,
    eps_min=None,
    revenue_min=None,
    net_income_min=None,
    roe_min=None,
    roa_min=None,
    profit_margin_min=None,
    institution_min=None,
    insider_min=None,
    close_min=None,
    close_max=None,
    distance_52w_high_max=None,
    distance_52w_low_max=None,
    volume_ratio_min=None,
):
    if q:
        pattern = f"%{q.strip()}%"
        query = query.filter(or_(Company.symbol.ilike(pattern), Company.name.ilike(pattern), Company.isin.ilike(pattern)))
    if sector:
        query = query.filter(Company.sector == sector)
    if industry:
        query = query.filter(Company.industry == industry)
    bounds = [
        ("market_cap", market_cap_min, market_cap_max),
        ("close", close_min, close_max),
    ]
    for key, lower, upper in bounds:
        expr = expressions[key]
        if lower is not None:
            query = query.filter(expr >= lower)
        if upper is not None:
            query = query.filter(expr <= upper)
    minimums = [
        ("trailing_eps", eps_min),
        ("revenue", revenue_min),
        ("net_income", net_income_min),
        ("return_on_equity", roe_min),
        ("return_on_assets", roa_min),
        ("profit_margin", profit_margin_min),
        ("institution_percent", institution_min),
        ("insider_percent", insider_min),
        ("volume_ratio", volume_ratio_min),
    ]
    for key, value in minimums:
        if value is not None:
            query = query.filter(expressions[key] >= value)
    if distance_52w_high_max is not None:
        query = query.filter(expressions["distance_52w_high"] <= distance_52w_high_max)
    if distance_52w_low_max is not None:
        query = query.filter(expressions["distance_52w_low"] <= distance_52w_low_max)
    return query


def _screener_row_dict(row):
    data = {}
    for key in SCREENER_COLUMN_LABELS:
        value = getattr(row, key, None)
        if isinstance(value, (datetime, date)):
            value = value.isoformat()
        elif isinstance(value, float) and not math.isfinite(value):
            value = None
        data[key] = value
    return data


def _screener_facets(db: Session, market: str):
    exchanges = _screener_market_exchanges(market)

    sector_query = (
        db.query(Company.sector)
        .filter(
            Company.is_active == 1,
            Company.exchange.in_(exchanges),
            Company.sector.isnot(None),
            Company.sector != "",
        )
    )
    sector_query = apply_eligible_equity_filter(sector_query, exchanges)
    sectors = [
        value for (value,) in sector_query.distinct().order_by(Company.sector.asc()).all() if value
    ]

    industry_query = (
        db.query(Company.industry)
        .filter(
            Company.is_active == 1,
            Company.exchange.in_(exchanges),
            Company.industry.isnot(None),
            Company.industry != "",
        )
    )
    industry_query = apply_eligible_equity_filter(industry_query, exchanges)
    industries = [
        value for (value,) in industry_query.distinct().order_by(Company.industry.asc()).all() if value
    ]
    return {"sectors": sectors[:500], "industries": industries[:1000]}


@router.get("/screener")
def get_universe_screener(
    market: str = "ALL",
    q: str | None = None,
    sector: str | None = None,
    industry: str | None = None,
    market_cap_min: float | None = None,
    market_cap_max: float | None = None,
    eps_min: float | None = None,
    revenue_min: float | None = None,
    net_income_min: float | None = None,
    roe_min: float | None = None,
    roa_min: float | None = None,
    profit_margin_min: float | None = None,
    institution_min: float | None = None,
    insider_min: float | None = None,
    close_min: float | None = None,
    close_max: float | None = None,
    distance_52w_high_max: float | None = None,
    distance_52w_low_max: float | None = None,
    volume_ratio_min: float | None = None,
    sort_by: str = "symbol",
    sort_dir: str = "asc",
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=10, le=100),
    db: Session = Depends(get_db),
):
    """Paginated multi-stock screener using only real stored DB data."""
    query, expressions = _screener_query_parts(db, market)
    query = _apply_screener_filters(
        query, expressions, q, sector, industry, market_cap_min, market_cap_max,
        eps_min, revenue_min, net_income_min, roe_min, roa_min, profit_margin_min, institution_min, insider_min, close_min,
        close_max, distance_52w_high_max, distance_52w_low_max, volume_ratio_min,
    )
    total = int(query.count())
    sort_expr = expressions.get(sort_by, Company.symbol)
    order = sort_expr.desc() if (sort_dir or "asc").lower() == "desc" else sort_expr.asc()
    rows = query.order_by(order.nullslast(), Company.symbol.asc()).offset((page - 1) * page_size).limit(page_size).all()
    return _json_safe({
        "market": market.upper(),
        "page": page,
        "page_size": page_size,
        "total": total,
        "pages": max(1, math.ceil(total / page_size)) if page_size else 1,
        "rows": [_screener_row_dict(row) for row in rows],
        "facets": _screener_facets(db, market),
        "coverage": universe_data_status(db, market),
        "columns": [{"key": key, "label": label} for key, label in SCREENER_COLUMN_LABELS.items()],
        "default_columns": SCREENER_DEFAULT_COLUMNS,
        "data_rule": "Stored provider/database values only. Missing values remain N/A and are never fabricated.",
    })


@router.get("/screener-export")
def export_universe_screener(
    market: str = "ALL",
    q: str | None = None,
    sector: str | None = None,
    industry: str | None = None,
    market_cap_min: float | None = None,
    market_cap_max: float | None = None,
    eps_min: float | None = None,
    revenue_min: float | None = None,
    net_income_min: float | None = None,
    roe_min: float | None = None,
    roa_min: float | None = None,
    profit_margin_min: float | None = None,
    institution_min: float | None = None,
    insider_min: float | None = None,
    close_min: float | None = None,
    close_max: float | None = None,
    distance_52w_high_max: float | None = None,
    distance_52w_low_max: float | None = None,
    volume_ratio_min: float | None = None,
    sort_by: str = "symbol",
    sort_dir: str = "asc",
    columns: str | None = None,
    db: Session = Depends(get_db),
):
    """Export the complete currently filtered screener result to Excel."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Excel export dependency is unavailable: {exc}")

    query, expressions = _screener_query_parts(db, market)
    query = _apply_screener_filters(
        query, expressions, q, sector, industry, market_cap_min, market_cap_max,
        eps_min, revenue_min, net_income_min, roe_min, roa_min, profit_margin_min, institution_min, insider_min, close_min,
        close_max, distance_52w_high_max, distance_52w_low_max, volume_ratio_min,
    )
    sort_expr = expressions.get(sort_by, Company.symbol)
    order = sort_expr.desc() if (sort_dir or "asc").lower() == "desc" else sort_expr.asc()
    rows = query.order_by(order.nullslast(), Company.symbol.asc()).limit(10000).all()

    requested = [c.strip() for c in (columns or "").split(",") if c.strip()]
    selected = [c for c in requested if c in SCREENER_COLUMN_LABELS] or list(SCREENER_DEFAULT_COLUMNS)
    if "symbol" not in selected:
        selected.insert(0, "symbol")

    wb = Workbook()
    ws = wb.active
    ws.title = "Filtered Stocks"
    ws.append([SCREENER_COLUMN_LABELS[c] for c in selected])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        item = _screener_row_dict(row)
        ws.append([item.get(c) for c in selected])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, key in enumerate(selected, start=1):
        max_len = len(SCREENER_COLUMN_LABELS[key])
        for cell in ws[get_column_letter(idx)][1: min(ws.max_row, 300)]:
            max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
        ws.column_dimensions[get_column_letter(idx)].width = min(32, max(12, max_len + 2))

    meta = wb.create_sheet("Filter Summary")
    meta.append(["Filter", "Value"])
    for c in meta[1]: c.font = Font(bold=True)
    for key, value in [
        ("Market", market), ("Search", q), ("Sector", sector), ("Industry", industry),
        ("Market Cap Min", market_cap_min), ("Market Cap Max", market_cap_max),
        ("EPS Min", eps_min), ("Revenue Min", revenue_min), ("Net Income Min", net_income_min),
        ("ROE Min", roe_min), ("ROA Min", roa_min), ("Profit Margin Min", profit_margin_min),
        ("Institution Min", institution_min), ("Insider Min", insider_min),
        ("Close Min", close_min), ("Close Max", close_max),
        ("Distance 52W High Max", distance_52w_high_max),
        ("Distance 52W Low Max", distance_52w_low_max),
        ("Volume Ratio Min", volume_ratio_min), ("Sort By", sort_by), ("Sort Direction", sort_dir),
        ("Exported Rows", len(rows)),
    ]:
        meta.append([key, value])
    meta.append(["Data Rule", "Stored provider/database values only; missing fields remain blank/N/A."])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"{market.upper()}_filtered_stock_screener.xlsx"
    payload = stream.getvalue()
    return Response(
        content=payload,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Content-Length": str(len(payload)),
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )
